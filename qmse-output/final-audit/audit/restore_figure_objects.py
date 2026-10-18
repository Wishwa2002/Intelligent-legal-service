from copy import deepcopy
from pathlib import Path
import re,json,hashlib
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
P=Path('qmse-output/final-audit');f=P/'final-report/SE3110_Testing_Report_GR42_Consolidated_Final.docx'
d=Document(f);o=Document(P/'backup/SE3110_Testing_Report_GR42_Historical_Verification.docx')
assert len(d.inline_shapes)==0
caps={int(re.match(r'Figure\s+6\.(\d+)',p.text).group(1)):p for p in d.paragraphs if re.match(r'Figure\s+6\.(\d+)',p.text)}
images={hashlib.sha256(rel.target_part.blob).hexdigest():rel.rId for rel in d.part.rels.values() if rel.reltype.endswith('/image')}
count=0
for i,p in enumerate(o.paragraphs):
 drawings=p._p.findall('.//'+qn('w:drawing'))
 if not drawings:continue
 preceding=[pp for pp in o.paragraphs[:i] if re.match(r'Figure\s+6\.(\d+)',pp.text)]
 no=int(re.match(r'Figure\s+6\.(\d+)',preceding[-1].text).group(1));cap=caps[no];target=cap._p.getnext()
 assert target.tag==qn('w:p')
 from docx.text.paragraph import Paragraph
 pp=Paragraph(target,cap._parent)
 for drawing in drawings:
  clone=deepcopy(drawing)
  for b in clone.iter(qn('a:blip')):
   oldrid=b.get(qn('r:embed'));sha=hashlib.sha256(o.part.related_parts[oldrid].blob).hexdigest();b.set(qn('r:embed'),images[sha])
  pp.add_run()._r.append(clone);count+=1
ci=json.loads((P/'historical-ci-index.json').read_text());cs={c['id']:c for c in ci['screenshots'] if 'image' in c}
for no,label in enumerate(['backend-success','frontend-success','mobile-retest-success','mobile-failure','frontend-error'],37):
 from docx.text.paragraph import Paragraph
 pp=Paragraph(caps[no]._p.getprevious(),caps[no]._parent);pp.add_run().add_picture(cs[label]['image'],width=Inches(6));count+=1
vm=json.loads((P/'security/visual-manifest.json').read_text())
for no,label in [(42,'before'),(43,'after')]:
 pp=Paragraph(caps[no]._p.getprevious(),caps[no]._parent);pp.add_run().add_picture(next(x['image'] for x in vm if x['label']==label),width=Inches(6));count+=1
assert len(d.inline_shapes)==43
# Preserve unique drawing object IDs while keeping image bytes and sizes intact.
for i,shape in enumerate(d.inline_shapes,1):shape._inline.docPr.set('id',str(i))
d.save(f)
(P/'audit/figure-restoration.json').write_text(json.dumps({'original_objects_restored':36,'new_genuine_objects_restored':7,'total':43,'source':'Preserved original report drawing XML and existing native CI/ZAP screenshots; no new image content generated'},indent=2))
print('Restored all 43 genuine figure objects')
