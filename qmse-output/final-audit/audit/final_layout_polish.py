from pathlib import Path
import json,hashlib
from docx import Document
from docx.shared import Pt,RGBColor
from docx.oxml.ns import qn
P=Path('qmse-output/final-audit'); f=P/'final-report/SE3110_Testing_Report_GR42_Consolidated_Final.docx'; d=Document(f)
# Record all previously viewed page images before final layout adjustments.
r=P/'final-report/render'
(P/'audit/visual-review-before-polish.json').write_text(json.dumps({'pages_individually_viewed':list(range(1,103)),'page_image_sha256':{str(i):hashlib.sha256((r/f'page-{i}.png').read_bytes()).hexdigest() for i in range(1,103)}},indent=2))
def setp(p,s):
 if p.runs:
  p.runs[0].text=s
  for x in p.runs[1:]: x.text=''
 else: p.add_run(s)
def cell(t,r,c,s): setp(d.tables[t].rows[r].cells[c].paragraphs[0],s)
cell(20,2,2,'Synthetic Customer JWT calls the current protected clerk-list route; related approval denials are separately verified.')
cell(20,2,4,'403 in the executed current-contract HTTP check; original manual session remains unverified.')
cell(20,2,5,'VERIFIED (current contract)')
cell(20,2,6,'TC-BK-05; historical-resolution evidence and final-audit/case-verification.json; Qmse_TC_C3_11 / Qmse_TC_C4_11')
cell(20,4,5,'BLOCKED (live acceptance)')
cell(20,5,4,'Earlier 15-request scan: 9 Low + 1 Informational. New same-six-path before/after scan: missing-nosniff instances 3 to 0; distinct scope.')
cell(20,5,6,'completion-20261008/evidence/zap-postgres.html; final-audit/security/zap-before/after-metadata.json; Figures 6.42–6.43')
new_decisions=[
'Existing 3–4,000-character contract retained. Eight DTO and eight browser boundary tests pass. Historical 10–1,000 expectation remains unapproved.',
'PROPOSED, disabled by default: 20% inclusive headroom signal. Demand 79% HEALTHY, 80% WATCH, 100% CONCERN. Existing 0.75 threshold is retained until approval; boundaries and defaults pass.',
'PROPOSED, disabled by default: 24-hour notice. Exact cutoff allowed; inside-window and past cancellation rejected without slot/history changes. Enabled and default-behaviour tests pass.',
'Orphan cleanup fix passes real PostgreSQL fault injection. Proposed, disabled 1–5-file service batch passes commit/rollback. No HTTP batch route; full historical case remains BLOCKED.'
]
for i,s in enumerate(new_decisions,1): cell(23,i,1,s)
findings=[
'Before: real PostgreSQL metadata-save failure left a physical file. After: same assertion passes. Proposed batch failure rolls back metadata, status and owned bytes in Development/Staging. Evidence: issue5-before.trx and issue5-retest.trx. Crash recovery and full historical HTTP atomicity remain unverified.',
'Low configuration finding corrected by middleware before error/auth responses. Four HTTP tests pass; same-six-path passive ZAP instances fall from 3 to 0. Evidence: security/zap-before/after-metadata.json and issue6-header.trx.',
'Non-placeholder historical database credential in original commit c0ea6458. Old value absent from current tracked contents; validity/rotation unverified. Owner must rotate/revoke privately; credential values are excluded.',
'Targeted Pillow, python-dotenv and python-multipart patches: 114 deterministic AI tests pass; pip check is clean. Returned advisory matches decrease from 70 to 21. Remaining framework/vector-store/test-tool findings need coordinated review.'
]
for i,s in enumerate(findings,1): cell(27,i,1,s)
cell(43,4,1,'Prior phase: 18 entries, including DEF-07 observation')
cell(43,4,3,'FA-DEF-01 / FA-SEC-01 are separate final-audit findings')
cell(44,11,1,'Earlier completion phase: upload retest 18/18; backend 326/328 with 2 skips')
cell(53,3,1,'Historical mismatches: TC-C103 fee update; TC-C106 reload expectation; TC-C107 10–1,000 length; TC-C112 20% headroom. Full live match/forecast cases TC-C104 / TC-C105 remain blocked. Current contract and disabled proposals are tested separately.')
cell(59,3,1,'TC-C2-01 historical Confirmed expectation differs from observed booking state. TC-C2-09 historical notice rule remains unresolved; proposed 24-hour policy is tested but disabled. TC-C2-05 full AI scheduling remains blocked.')
cell(65,3,1,'Original six failures: four accepted current-contract replacements pass (C3-02/03/05/06). C3-01 service/checklist identity and C3-04 Approved transition remain failed historical expectations. C3-08 OCR and C3-12 full atomic workflow remain blocked.')
cell(73,4,1,'Original spoof failure/retest, dated 607-pass baseline and 23-test extension (12 pass/11 fail) are preserved. Final 77-case ledger is 49 PASS / 16 FAIL / 12 BLOCKED. New regression populations are separately evidence-linked. All 36 existing figures retained; five CI and two ZAP figures added, for 43. Personal member assessment remains pending.')
appendix={
5:'Final 77-case ledger: 49 PASS, 16 FAIL, 12 BLOCKED; zero historical/unreviewed or not-executed. Earlier 35/30/12 ledger retained. Fourteen executed accepted replacements pass; blocked cases remain unfinished. See final-audit/index.md.',
7:'Earlier real PostgreSQL k6 and expanded passive ZAP reports preserved. Final six-path ZAP retest verifies nosniff remediation, with limited scan scope. Live AI latency acceptance and production capacity remain unverified.',
9:'Latest deterministic/mocked AI suite: 114 PASS after dependency patches. Earlier runs retained. Two prior live requests returned 504/503; no new live call or full expert-labelled Golden Case acceptance. See provider-diagnosis.md.',
10:'Final case statuses and actual results reconciled to evidence: 49 PASS / 16 FAIL / 12 BLOCKED. Historical observations and original expectations remain visible.',
11:'Four prior-phase fixes retained. Final-audit filesystem orphan and missing-nosniff fixes have before/after evidence. Historical mismatches are not automatically product defects.',
13:'Separate populations: dated 607 PASS / 5 SKIP baseline; 23-test historical extension (12 PASS / 11 FAIL); final backend 362 PASS / 1 FAIL / 2 SKIP. Other latest suites are listed separately in Section 4.',
14:'Published baseline 05359214df52fa743b234e1d77e6671f57c14dea preserved. Latest 33-file local overlay is uncommitted/unpublished, supplied as tested-source.patch and source-snapshot/. No final-overlay CI claim.',
16:'Current reproduction instructions: final-audit/reproduction.md; source-manifest.json identifies the unpublished overlay. Fresh isolated database, synthetic settings and dependencies required; no private credentials are packaged.',
17:'Pending member confirmation: personal execution, identity mapping, understanding and demonstration. Member-specific commands and questions are in final-audit/member-viva-guide.md.',
18:'Updated CLEAR disclosure preserves historical AI use and adds Codex-assisted fixes, test corrections, regression execution, CI evidence and document work. No student authorship inferred.'}
for row,s in appendix.items(): cell(74,row,2,s)
# Body rows of newly inserted tables must not inherit header shading/font.
for ti in (23,27,53,59,65,71):
 for ri,row in enumerate(d.tables[ti].rows[1:],1):
  for ci,c in enumerate(row.cells):
   for sh in c._tc.get_or_add_tcPr().findall(qn('w:shd')): sh.set(qn('w:fill'),'F2F4F8' if ri%2==0 else 'FFFFFF')
   for p in c.paragraphs:
    p.paragraph_format.keep_with_next=False
    for run in p.runs: run.font.color.rgb=RGBColor.from_string('111111');run.font.size=Pt(10);run.bold=(ci==0)
 # Remove only explicit page break inserted immediately before the new member evidence table.
 if ti in (53,59,65,71):
  prev=d.tables[ti]._element.getprevious()
  if prev is not None and prev.tag==qn('w:p') and not ''.join(prev.itertext()).strip():
   for br in list(prev.iter(qn('w:br'))):
    if br.get(qn('w:type'))=='page': br.getparent().remove(br)
replacements={
'12 component-assigned cases: 8 VERIFIED/PASS, 3 VERIFIED/FAIL, 1 BLOCKED':'12 component-assigned cases: 9 VERIFIED/PASS, 2 VERIFIED/FAIL, 1 BLOCKED',
'12 component-assigned cases: 4 VERIFIED/PASS, 6 VERIFIED/FAIL, 2 BLOCKED':'12 component-assigned cases: 8 VERIFIED/PASS, 2 VERIFIED/FAIL, 2 BLOCKED',
'Figure 6.20. Final xUnit execution on the tested source':'Figure 6.20. Earlier completion-phase xUnit execution on its tested source',
'Nine Low missing-nosniff instances and one Informational alert remain.':'This earlier scan recorded nine Low missing-nosniff instances and one Informational alert. Figures 6.42–6.43 show the later scoped header fix and retest.',
'final-audit/index.':'final-audit/index.md.',
'correction8a9b8a24':'correction 8a9b8a24','lines11 and30':'lines 11 and 30','SHAfb471202':'SHA fb471202','plus exit2':'plus exit 2',
'ZAP2.17.0':'ZAP 2.17.0','unpublishedoverlay overSHA05359214':'unpublished overlay over SHA 05359214',
'Before:3 alert instances of1':'Before: 3 alert instances of 1','source:final-audit':'source: final-audit',
'After:zeroalert instances; allsix responses havenosniff.':'After: zero alert instances; all six responses have nosniff.',
'This islimitedsample evidence, notwhole-application security.':'This is limited-sample evidence; it does not certify whole-application security.'}
for p in d.paragraphs:
 s=p.text
 for a,b in replacements.items(): s=s.replace(a,b)
 if s!=p.text: setp(p,s)
 if p.text.startswith('Cold PostgreSQLTLS'):
  setp(p,'Cold PostgreSQL TLS connection remains a historical discrepancy: an earlier final-audit run measured 585.067 ms against <200 ms. The latest full-run assertion passed, as did an earlier warm run; this does not demonstrate a stable cold-connection fix. API p95 is a separate measurement; dated k6 reports are unchanged.')
 if p.text.startswith('Earlier ZAP 2.17.0 passive15'):
  setp(p,'Earlier 15-request ZAP 2.17.0 evidence is preserved. Final same-six-path passive scans verify missing-nosniff remediation: 3 alert instances of one Low type before, 0 after. All six success/401/404 responses carry nosniff. Native HTML/XML/JSON and Figures 6.42–6.43 are retained. No active scan, spider or production target was used; scope is limited.')
 if p.text.startswith('Security incident:'):
  setp(p,'Security incident: non-placeholder historical PostgreSQL/Neon credential material was found in original commit c0ea6458. Its old value is absent from current tracked contents; validity and rotation are unknown. The owner must rotate/revoke it privately and review access. Targeted dependency patches reduced returned advisory matches from 70 across 6 packages to 21 across 3; aliases may duplicate findings. Pip check is clean and 114 deterministic AI tests pass. Starlette, ChromaDB and pytest need coordinated review. The npm advisory query remains blocked by TLS certificate-chain errors; historical Dependabot alerts do not certify this overlay. See final-audit/security/security-review.md/.json.')
 if p.text.startswith('New real PostgreSQLdevelopment'):
  setp(p,'Real PostgreSQL fault injection reproduced a Development-mode file orphan after metadata-save failure. Cleanup now removes owned bytes and leaves status PENDING with no file rows on the matching retest. The proposed, disabled batch service passes second-file validation/persistence rollback and two-file commit in Development/Staging. It operates on an existing request; no HTTP batch route or full request-plus-three-files atomicity is claimed. Filesystem/database crash recovery remains a limitation.')
 for comp,counts,body in [
 ('C1', '7 PASS / 4 FAIL / 2 BLOCKED','The existing 3–4,000-character query contract is retained and boundary-tested; the client upper guard is explicit. The optional inclusive 20% headroom proposal passes boundaries but is disabled; default 0.75 WATCH rules remain. Financial/catalogue expectations and live scoring/forecast remain unresolved.'),
 ('C2','9 PASS / 2 FAIL / 1 BLOCKED','Current-contract cancellation (POST/200/Cancelled with slot release) passes. Proposed 24-hour notice is disabled by default; enabled tests allow 1,440/1,441 minutes and reject 1,439/0/-1 without mutation. API capability and Flutter button agree. Historical notice/Confirmed expectations remain failed; full scheduling AI acceptance is blocked.'),
 ('C3','8 PASS / 2 FAIL / 2 BLOCKED','Four accepted replacements pass; service identity and Approved mapping remain historical failures. Spoofed-PDF regression evidence is preserved. Development file-orphan cleanup is fixed/retested on real PostgreSQL. Proposed service-batch rollback/commit passes but remains disabled without an HTTP/UI batch route. Full C3-12 atomic workflow and live OCR remain blocked.'),
 ('C4','8 PASS / 0 FAIL / 4 BLOCKED','JWT/RBAC/service-request evidence is retained. Missing-nosniff configuration is fixed and passively retested; historical credential exposure and remaining dependency advisories require owner action. Provider 503/504/timeout/schema resilience mocks pass; live consultation/relevance/disclaimer acceptance remains incomplete.')]:
  if p.text.startswith(comp+' remains'): setp(p,f'{comp}: {counts}. {body} Personal verification is pending member confirmation.')
 if p.text.startswith('Shared primary/support proposal:'):
  setp(p,'Shared primary/support proposal: Jeyashankar 6 catalogue/model cases; Buddhika 7 scheduling/mobile-intake cases; Vithusan 8 database/document cases, including already-verified DB-04; Wijayananda 7 authentication/client-security cases. All four support integrated E2E. The 28-case allocation is responsibility planning, not personal effort attribution; the earlier 27-case historical subset is preserved. See final-audit/individual-contribution-matrix.md and member-evidence-mapping.md.')
 if p.text.startswith('Approve orreject'):
  setp(p,'Approve or reject proposals before enabling defaults: 24-hour cancellation notice with exact cutoff allowed; inclusive 20% headroom signal; owned atomic batch on an existing request. Confirm financial, service/status, JWT, BCrypt and encrypted-storage decisions separately. The existing 3–4,000-character query contract is retained. Technical cleanup and security headers were authorised fixes, not claimed as prior approved SE3090 rules. See final-audit/requirements-decision-register.md.')
 if p.text.startswith('Address the verified Low missing nosniff'):
  setp(p,'Review and publish the locally verified nosniff correction only through an approved deployment change. Repeat representative PostgreSQL load after staging approval; current two-user read tests do not certify production capacity.')
 if p.text.startswith('Ownerpriority:'):
  setp(p,'Owner priority: revoke/rotate the historical database credential privately and assess access history. Review remaining Starlette/ChromaDB/pytest advisories and repair trusted npm CA configuration for advisory verification. Obtain expert Golden Case labels, provider availability and a new live-call budget. Decide Deed/Approved mapping and the historical atomic HTTP boundary. Each member confirms personal execution, account mapping and AI assistance. Historical CI logs are retained; final-branch CI needs separate publication/workflow approval.')
 if p.text.startswith('Codex-assisted work includes historicalverification'):
  setp(p,'Codex-assisted work includes historical verification, current-contract test corrections, isolated configuration, filesystem/security-header/client-capability fixes, disabled policy proposals, targeted dependency patches, regression execution, approved read-only CI log retrieval, genuine screenshots and report consolidation. Members must confirm their understanding and personal execution. Component assignment does not establish student authorship of these changes.')
 if p.text.startswith('Reproduction: clone officialpublishedbranch'):
  setp(p,'Clone the official published branch and verify SHA 05359214. Apply final-audit/tested-source.patch and copy the untracked files from source-snapshot/, then use a fresh owned loopback PostgreSQL cluster and synthetic accounts. Exact versions, commands and disabled flags are in final-audit/reproduction.md and source-manifest.json. This overlay is uncommitted/unpublished. Never use the fixtures on production, call live Gemini without a new budget or overwrite delivered logs.')
# Clarify earlier mobile result without removing its evidence.
for row in d.tables[76].rows:
 for c in row.cells:
  for p in c.paragraphs:
   if 'Recorded final 49-pass output:' in p.text:
    setp(p,p.text.replace('Recorded final 49-pass output:','Earlier completion-phase 49-pass output:')+' Latest separate full Flutter run: 60 PASS; final-audit/debug/evidence/issue5-flutter-final.log.')
d.save(f)
print('Saved targeted result, provenance and formatting corrections; 43 figures retained.')
