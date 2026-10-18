from pathlib import Path
from docx import Document
P=Path('qmse-output/final-audit');f=P/'final-report/SE3110_Testing_Report_GR42_Consolidated_Final.docx';d=Document(f)
def setp(p,s):
 if p.runs:
  p.runs[0].text=s
  for r in p.runs[1:]:r.text=''
 else:p.add_run(s)
for p in d.paragraphs:
 if p.text.startswith('Final failure classification (original 30 FAIL cases):'):
  setp(p,'Original 30 failures: 0 established approved-requirement product defects; 18 historical expectation mismatches; 10 obsolete scenarios; 0 established setup defects; 2 undetermined. Fourteen accepted replacements pass; 16 failures remain. Separately, this audit fixed the reproduced filesystem orphan defect (FA-DEF-01) and missing-nosniff configuration (FA-SEC-01). Original observations remain preserved. See final-audit/failure-classification.md and requirements-decision-register.md.')
 if p.text.startswith('Final audit: no live provider requests.'):
  setp(p,'Final audit: no live provider requests. Configured model: gemini-3.8-flash; account/model availability remains unverified. Prior 503/504 causes cannot be narrowed beyond redacted evidence. Real service startup and internal authentication were verified; empty-key generation safely returns 503. The latest deterministic run passes 114 tests, including six provider-resilience cases, after dependency patches. Earlier 108-test output is retained. Raw retrieval is not generated legal advice or disclaimer evidence. Expert labels and universal-disclaimer acceptance remain blocked. See final-audit/provider-diagnosis.md and golden-case-evaluation-template.json.')
 if p.text.startswith('Final compliance position:'):
  setp(p,'Final compliance: 77 cases comprise 49 PASS, 16 FAIL and 12 BLOCKED. Fourteen accepted replacements executed; no complete blocked case was promoted. Earlier 607 machine passes and the 23-test extension remain separate dated populations. Five historical CI figures and two local ZAP figures were added; the published QMSE SHA has zero observed CI runs. Member demonstrations and proposed business rules await confirmation. See final-audit/execution-summary.json and member-confirmation-checklist.md.')
c=d.tables[6].rows[3].cells[1]
setp(c.paragraphs[0],'Earlier integration: API 55443, FastAPI 55444, React 55447, Flutter web 55448, ZAP 55449. Final local audit: API 55543, FastAPI 55544, isolated PostgreSQL 55546, React 55547, Flutter web 55548, ZAP 55549. These are separate dated environments; all phase-owned services were stopped after testing.')
c=d.tables[6].rows[6].cells[1]
setp(c.paragraphs[0],'Published group branch se3110-qmse-testing remains at 05359214df52fa743b234e1d77e6671f57c14dea. Local branch: docs/se3110-qmse-testing. Final regressions test this baseline plus the 33-file unpublished AI-assisted overlay recorded in final-audit/source-manifest.json and source-snapshot/. Final source differs from the published baseline; no CI result certifies the overlay. No new commit or push was performed.')
d.save(f)
print('Targeted semantic corrections saved')
