from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import re,json
P=Path('qmse-output/final-audit');path=P/'final-report/SE3110_Testing_Report_GR42_Consolidated_Final.docx';d=Document(path);old=Document('qmse-output/historical-resolution/final-report/SE3110_Testing_Report_GR42_Historical_Verification.docx');text='\n'.join(p.text for p in old.paragraphs)+'\n'+'\n'.join(c.text for t in old.tables for r in t.rows for c in r.cells)
maps={'05362214':'05359214','0536221':'0535921'}
for h in re.findall(r'\b[0-9a-f]{7,40}\b',text):
 mutated=h.replace('363','365').replace('359','362')
 if mutated!=h:maps[mutated]=h
count=0
for el in d.element.iter(qn('w:t')):
 s=el.text or '';v=s
 for a,b in sorted(maps.items(),key=lambda x:-len(x[0])):v=v.replace(a,b)
 if v!=s:el.text=v;count+=1
for p in d.paragraphs:
 if p.text.startswith('Consolidated 77-case outcome:'):
  for r in p.runs:r.text=''
  p.runs[0].text='Consolidated 77-case outcome: 49 PASS, 16 FAIL and 12 BLOCKED. Fourteen executed replacements use the explicitly accepted current contract; original failure artifacts remain. The development upload orphan and missing-nosniff configuration were fixed and retested. Cancellation, headroom and atomic-batch proposals remain disabled pending group approval. Latest separate full regressions: backend365 total (362 PASS,1 FAIL,2 SKIPPED), Flutter60 PASS, deterministicAI114 PASS, Node19 PASS, baselinePlaywright111 PASS/3 SKIPPED. Earlier607 passes and23 native results remain distinct dated populations. Twelve historical CI logs are now retained; exact published SHA05359214 has zero observed runs. No push, deployment or new live provider call.'
d.save(path);(P/'audit/source-reference-correction.json').write_text(json.dumps({'corrected_text_nodes':count,'mapping':maps,'actual_head':'05359214df52fa743b234e1d77e6671f57c14dea','reason':'Correct collateral source-identity substitution before delivery; no Git data changed'},indent=2));print('Corrected source references',count)
