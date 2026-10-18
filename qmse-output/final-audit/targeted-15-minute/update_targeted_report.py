from pathlib import Path
from docx import Document
import shutil,json,hashlib
P=Path('qmse-output/final-audit');O=P/'targeted-15-minute';F=P/'final-report';name='SE3110_Testing_Report_GR42_Consolidated_Final';B=O/'report-before';B.mkdir(exist_ok=True)
for ext in ('docx','pdf'):
 s=F/(name+'.'+ext);shutil.copy2(s,B/s.name)
(O/'report-input-hashes.json').write_text(json.dumps({p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in B.iterdir()},indent=2))
d=Document(F/(name+'.docx'))
p=d.paragraphs[68]
p.text+=' Targeted follow-up: 15 authentication tests executed on 8 October 2026, 14 PASS / 1 FAIL / 0 SKIPPED. Three new current-contract hashing regressions pass; the original cost-12 assertion still fails. Approval remains unconfirmed, so cost 11 is retained. See final-audit/targeted-15-minute/focused-authentication.trx and PasswordCurrentContractTests.cs; this focused population is not added to the full-suite totals.'
p=d.paragraphs[344]
p.text=p.text.replace('validity and rotation are unknown.', 'validity and rotation remain UNCONFIRMED; no owner/provider revocation evidence was supplied. A targeted 8 October 2026 comparison found no historical credential value or high-confidence token/private-key pattern in PDF extracted text or DOCX text. Image-only evidence was previously visually reviewed; OCR was unavailable. See final-audit/targeted-15-minute/credential-pdf-check.json and owner-action-warning.md.')
p=d.paragraphs[438]
p.text=p.text.replace('No push, deployment or new live provider call.', 'Focused authentication follow-up: 14 PASS / 1 FAIL / 0 SKIPPED, separately recorded; cost 11 remains pending approval of cost 12. Credential rotation remains unconfirmed. No push, deployment or new live provider call.')
for t in d.element.xpath('.//w:t'):
 if t.text:
  t.text=t.text.replace('Latest 33-file local overlay','Prior 33-file local overlay').replace('plus the 33-file unpublished AI-assisted overlay','plus the prior 33-file unpublished AI-assisted overlay')
d.save(F/(name+'.docx'))
print('Targeted results, security note and conclusion updated; figures:',len(d.inline_shapes))
