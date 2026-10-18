from pathlib import Path
import json,re,zipfile,hashlib,collections,xml.etree.ElementTree as ET,shutil,subprocess
from datetime import datetime,timezone
import pymupdf
from docx import Document
R=Path.cwd();O=R/'qmse-output/historical-resolution';F=O/'final-report';name='SE3110_Testing_Report_GR42_Historical_Verification';shutil.copy2(F/'render'/(name+'.pdf'),F/(name+'.pdf'));D=Document(F/(name+'.docx'));old=Document(R/'qmse-output/final-report/SE3110_Testing_Report_GR42_Final.docx');pdf=pymupdf.open(F/(name+'.pdf'));checks={}
def check(k,v):
 checks[k]=bool(v)
 if not v:raise AssertionError(k)
check('docx_opens',zipfile.ZipFile(F/(name+'.docx')).testzip() is None);check('pdf_opens',len(pdf)>0);check('all_71_tables_3_sections',len(D.tables)==len(old.tables)==71 and len(D.sections)==len(old.sections)==3)
check('section_geometry_unchanged',[s._sectPr.xml for s in D.sections]==[s._sectPr.xml for s in old.sections]);check('cover_students_preserved',[[c.text for c in rr.cells] for rr in D.tables[1].rows]==[[c.text for c in rr.cells] for rr in old.tables[1].rows])
for ti in range(10,19):
 check('case_design_preserved_table_'+str(ti),[[c.text for c in rr.cells[:5]] for rr in D.tables[ti].rows]==[[c.text for c in rr.cells[:5]] for rr in old.tables[ti].rows])
def media(p):
 with zipfile.ZipFile(p) as z:return {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
check('all_30_original_images_preserved',media(R/'qmse-output/final-report/SE3110_Testing_Report_GR42_Final.docx')<=media(F/(name+'.docx')));check('36_docx_figures',len(D.inline_shapes)==36)
cs=json.loads((O/'test-results/case-verification.json').read_text());n=json.loads((O/'test-results/new-case-verification.json').read_text());s=json.loads((O/'test-results/summary.json').read_text());check('77_distinct_case_ids',len(cs)==len({c['test_id'] for c in cs})==77);check('35_pass_30_fail_12_block',dict(collections.Counter(c['status'] for c in cs))=={'VERIFIED/PASS':35,'VERIFIED/FAIL':30,'BLOCKED':12});check('49_reviewed_39_verified_10_blocked',len(n)==49 and sum(c['status'].startswith('VERIFIED/') for c in n)==39 and sum(c['status']=='BLOCKED' for c in n)==10)
for c in cs:
 for e in c['evidence']:check('case_evidence_exists_'+e,(R/e).exists())
for c in n:check('blocked_no_execution_timestamp_'+c['test_id'],c['execution_date'] is None if c['status']=='BLOCKED' else bool(c['execution_date']))
case_rows={rr.cells[0].text.strip():rr for t in D.tables[10:19] for rr in t.rows[1:]}
check('77_report_rows_match_json',all(case_rows[c['test_id']].cells[6].text==c['status'] for c in cs))
ns={'t':'http://microsoft.com/schemas/VisualStudio/TeamTest/2010'};native=[]
for name_trx in ['password-cost.trx','postgres-setup.trx','ef-duplicate.trx']:
 a=ET.parse(O/'evidence'/name_trx).find('.//t:Counters',ns).attrib;native.append({'report':name_trx,**{k:int(a[k]) for k in ['total','passed','failed']}})
check('new_xunit_4_2_2',tuple(sum(x[k] for x in native) for k in ['total','passed','failed'])==(4,2,2));check('prior_functional_hash_pass_retained',ET.parse(O/'evidence/password.trx').find('.//t:Counters',ns).attrib['passed']=='1')
ai=ET.parse(O/'evidence/ai-retrieval.xml').find('testsuite').attrib;check('new_pytest_5_3_2',ai['tests']=='5' and ai['failures']=='2' and ai['errors']=='0')
fl=[]
for line in (O/'evidence/flutter-historical-final.log').read_text().splitlines():
 try:x=json.loads(line)
 except:continue
 if isinstance(x,dict) and x.get('type')=='testDone' and not x.get('hidden'):fl.append(x)
check('new_flutter_5_2_3',len(fl)==5 and sum(x['result']=='success' for x in fl)==2)
pw=json.loads((O/'evidence/playwright-historical.json').read_text())['stats'];check('new_playwright_9_5_4',pw['expected']==5 and pw['unexpected']==4 and pw['skipped']==0);ret=json.loads((O/'evidence/playwright-no-slot.json').read_text())['stats'];check('semantic_no_slot_retest_fail',ret['unexpected']==1)
base=ET.parse(R/'qmse-output/completion-20261008/evidence/backend-final-verified.trx').find('.//t:Counters',ns).attrib;check('baseline_xunit_326_pass_2_skip',base['total']=='328' and base['passed']=='326' and base['failed']=='0');basepw=json.loads((R/'qmse-output/completion-20261008/evidence/playwright-final.json').read_text())['stats'];check('baseline_playwright_111_3_skip',basepw['expected']==111 and basepw['unexpected']==0 and basepw['skipped']==3);baseai=ET.parse(R/'qmse-output/completion-20261008/evidence/ai-completion.xml').find('testsuite').attrib;check('baseline_pytest_102',baseai['tests']=='102' and baseai['failures']=='0');check('baseline_607_unchanged',s['baseline_machine_passes_preserved']==607 and 326+19+49+111+102==607)
norm=lambda s:re.sub(r'\s+',' ',s).strip();texts=[norm(p.get_text()) for p in pdf];full='\n'.join(texts)
for im in json.loads((O/'evidence/figure-index.json').read_text()):
 found=[i+1 for i,t in enumerate(texts) if norm('Figure '+im['figure']+'. '+im['title']) in t];check('new_figure_one_caption_'+im['figure'],len(found)==1);check('new_figure_same_page_image_'+im['figure'],len(pdf[found[0]-1].get_image_info())>0);check('new_figure_source_exists_'+im['figure'],all((R/'qmse-output'/p).exists() for p in [im['image']]+im['source']))
check('36_pdf_images',sum(len(p.get_image_info()) for p in pdf)==36);check('no_blank_pages',all(len(re.sub(r'SE3110 Software Testing Report.*|Page \d+','',p.get_text()).strip())>40 for p in pdf));check('page_numbers',all('Page '+str(i+1) in p.get_text() for i,p in enumerate(pdf)));check('image_bounds',all(p.rect.contains(pymupdf.Rect(im['bbox'])) for p in pdf for im in p.get_image_info()));check('no_text_off_page',all(p.rect.contains(pymupdf.Rect(w[:4])) for p in pdf for w in p.get_text('words')))
check('all_member_sections',all(s in full for s in ['9.1 Jeyashankar','9.2 Buddhika','9.3 Vithusan','9.4 Wijayananda']))
for bm,t in json.load(open(O/'audit/toc-pages.json')).items():check('toc_target_'+bm,norm(t['title']) in texts[t['page']-1])
for p,h in json.load(open(O/'audit/input-baseline.json'))['inputs'].items():check('original_unchanged_'+Path(p).name,hashlib.sha256(Path(p).read_bytes()).hexdigest()==h)
check('app_tracked_source_unchanged',subprocess.check_output(['git','diff','--name-only','HEAD'],text=True).strip()=='');check('approved_baseline_head_preserved',subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip()==s['published_sha']);check('owned_services_stopped',json.loads((O/'audit/cleanup.json').read_text())['confirmed_stopped'])
check('new_cases_all_ai_assisted',all('AI-assisted' in c['authorship'] for c in n));check('live_probe_503_two_calls',json.loads((O/'evidence/live-gemini.json').read_text())['cumulative_attempted_calls']==2 and json.loads((O/'evidence/live-gemini.json').read_text())['results'][0]['code']==503)
record={'checked_utc':datetime.now(timezone.utc).isoformat(),'pdf_pages':len(pdf),'figures':36,'checks':checks,'native_xunit':native,'new_case_summary':s,'original_document_preserved':True};(O/'audit/final-quality-check.json').write_text(json.dumps(record,indent=2));print({'passed_checks':sum(checks.values()),'pdf_pages':len(pdf),'figures':36})
