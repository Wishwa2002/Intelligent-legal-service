from pathlib import Path
from copy import deepcopy
import json,re,shutil,hashlib,collections
from docx import Document
from docx.shared import Pt,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path.cwd();O=R/'qmse-output/historical-resolution';F=O/'final-report';F.mkdir(exist_ok=True);source=R/'qmse-output/final-report/SE3110_Testing_Report_GR42_Final.docx';name='SE3110_Testing_Report_GR42_Historical_Verification';D=Document(source);P=list(D.paragraphs);T=D.tables;cases=json.loads((O/'test-results/case-verification.json').read_text());cs={x['test_id']:x for x in cases};new=json.loads((O/'test-results/new-case-verification.json').read_text());S=json.loads((O/'test-results/summary.json').read_text());edits=[]
def put(p,s):
 old=p.text
 if old==s:return
 fmt=deepcopy(p.runs[0]._r.find(qn('w:rPr'))) if p.runs and p.runs[0]._r.find(qn('w:rPr')) is not None else None
 for c in list(p._p):
  if c.tag!=qn('w:pPr'):p._p.remove(c)
 run=p.add_run(str(s))
 if fmt is not None:run._r.insert(0,fmt)
 edits.append({'before':old,'after':s})
def cell(c,s):
 put(c.paragraphs[0],s)
 for p in c.paragraphs[1:]:c._tc.remove(p._p)
def row(t,i,v):
 for c,s in zip(t.rows[i].cells,v):cell(c,str(s))
def addrow(t,v):
 from docx.table import _Row
 e=deepcopy(t.rows[-1]._tr);t._tbl.append(e);rr=_Row(e,t)
 for c,s in zip(rr.cells,v):cell(c,str(s))
def before(p,text,size=10):
 x=p.insert_paragraph_before(text);x.paragraph_format.space_after=Pt(6)
 for r in x.runs:r.font.size=Pt(size)
 return x
short={
'TC-BK-02':'Actual seven one-hour slots start at 09:00, 11:00, 12:00, 13:00, 14:00, 15:00 and 16:00. The reserved 10:00 slot is excluded; synthetic lawyer 1 maps to a GUID.',
'TC-BK-03':'Historical BarNumber is not the validated field. Empty BarNumber plus valid current license yields 201; empty current license control yields 400.',
'TC-BK-04':'Login HTTP 200. JWT uses nameidentifier/emailaddress/role URI claim names; historical sub/plain email claims absent.',
'TC-BK-06':'Historical specializationId query returns all five seeded lawyers. Current specialization parameter returns the three Corporate lawyers.',
'TC-C103':'PUT HTTP 200, but Price is ignored and actual LegalServices row has no fee/Price column; 25000 not persisted.',
'TC-C2-01':'Actual booking HTTP 201, status Requested; historical immediate Confirmed expectation differs.',
'TC-C2-02':'Second real PostgreSQL-backed booking for the same lawyer/slot HTTP 409; original row retained.',
'TC-C2-03':'Historical PUT cancel returns 405. Current POST cancel control returns 200 and Cancelled after confirming the booking.',
'TC-C3-01':'Exact current service 1 is Contract Review & Amendment, with three required documents. Request HTTP 201/PENDING; historical four-item Deed mapping differs.',
'TC-C3-02':'Historical POST /api/document-files returns 404. Actual per-request upload of genuine 2 MiB PDF returns 201 and is persisted.',
'TC-C3-03':'Historical route returns 404. Current per-request route rejects genuine 14 MiB PDF with HTTP 400, without an extra file record.',
'TC-C3-04':'Historical clerk approval route returns 404. No Approved transition was invented.',
'TC-C3-06':'Historical /department subroute returns 404. Current PUT /api/clerks/3 control updates Property department with Contact supplied.',
'TC-E2E-02':'Actual documentation creation HTTP 201 with requestId, but no historical tracking-number field.',
'TC-E2E-03':'Actual isolated PostgreSQL row persisted with PENDING, not historical PendingChecklist. No production Neon access.',
'TC-E2E-04':'Actual FastAPI historical /api/ai/document/audit-checklist POST returns 404; current agent-session routes are separate contracts.',
'TC-DB-03':'All 15 actual migrations applied. Actual schema has '+str(next(x for x in new if x['test_id']=='TC-DB-03')['observed']['tables'])+' public tables including migration history, not historical 11.',
'TC-FE-01':'Actual React catalogue renders Kavinda Silva; no historical fee display/Book control. Explicit UI API fixture, not live catalogue proof.',
'TC-C106':'Actual React catalogue renders lawyer but contains no historical card availability switch. Current date filter/Edit/Delete differ.',
'TC-FE-02':'Actual booking approval remains disabled without a slot; no inline slot/date error alert appears. Instructional hints differ from historical error design. UI fixture only.',
'TC-FE-03':'Actual unauthenticated clerk route redirects to /login; historical redirect query parameter is absent.',
'TC-FE-04':'Actual Axios module adds the expected Bearer header to intercepted service-request call. Synthetic token; header injection only.',
'TC-FE-05':'Actual React HTTP-500 fixture displays a usable alert and clickable Retry; no page crash. Backend availability is not evaluated.',
'TC-C2-04':'Actual POST careers HTTP 201; Senior Counsel visible on real React public careers page using PostgreSQL data.',
'TC-C2-06':'Actual select-lawyer/date/slot/approval UI displays Appointment Created. Current approval label differs; API responses explicitly mocked.',
'TC-C4-06':'Actual React login with isolated API stores token; Logout clears it and returns to /login. Credentials/JWT excluded.'}
# Existing 77 case designs preserved; actual/status/evidence cells updated, expected/input/precondition cells unchanged.
for ti in range(10,19):
 for rr in T[ti].rows[1:]:
  id=rr.cells[0].text.strip()
  if id not in cs:continue
  x=cs[id];cell(rr.cells[6],x['status'])
  if id not in {z['test_id'] for z in new}:continue
  obs=short.get(id,x['observed'])
  if not isinstance(obs,str):obs=json.dumps(obs)
  if len(obs)>620:obs=obs[:615]+'…'
  date=(x.get('execution_date') or x['investigation_date']).split('T')[0]
  cell(rr.cells[5],obs+(' Investigation (full case not executed): ' if x['status']=='BLOCKED' else ' Execution: ')+date+' (UTC).')
  primary=x['evidence'][0].replace('qmse-output/','');cell(rr.cells[7],primary+'; historical-resolution/test-results/case-verification.json. AI-assisted replay; full scope/inputs recorded.')
put(P[12],'Historical verification revision: 8 October 2026. The published application baseline remains 05359214df52fa743b234e1d77e6671f57c14dea on Wishwa2002/Intelligent-legal-service/se3110-qmse-testing. New AI-assisted acceptance replays are a separately delivered, uncommitted local test overlay; no push, deployment, business-rule change or production access occurred. The previous 78-page report and all native evidence remain unchanged.')
put(P[41],'VERIFIED/PASS means an executed scoped scenario matched its recorded expectation; VERIFIED/FAIL means an observed mismatch, including historical route, precondition or requirement conflicts. A report-expectation failure is not automatically an approved product defect. BLOCKED means full acceptance could not run; any executed subcriteria are explicitly recorded. HISTORICAL/UNVERIFIED remains available for uninvestigated claims. Exact current/historical inputs, dates and evidence are in historical-resolution/test-results/case-verification.json; synthetic fixture dates are not execution dates.')
before(P[42],'All 49 historical cases were inventoried before execution. This revision obtains 39 scoped results (12 PASS, 27 FAIL) and investigates 10 blocked cases. The full original 77-case set now has 35 VERIFIED/PASS, 30 VERIFIED/FAIL and 12 BLOCKED; HISTORICAL/UNVERIFIED and NOT EXECUTED are zero. Reclassification to BLOCKED does not mean completion.',10)
cell(T[6].cell(3,1),T[6].cell(3,1).text+' New historical replay: API 55543; FastAPI 55544; React 55547; Flutter web 55548; separate isolated PostgreSQL 55546.')
cell(T[6].cell(4,1),T[6].cell(4,1).text+' Historical replay creates a separate TLS-enabled qmse_history database; 15 actual migrations applied.')
cell(T[6].cell(6,1),T[6].cell(6,1).text+' New tests are local uncommitted AI-assisted additions and are supplied as hashed source snapshots.')
put(P[60],'The earlier baseline contains 607 passing machine cases and five skips; its original native logs are preserved. New distinct extensions executed 23 native cases: 12 PASS, 11 FAIL, no skips. They include setup, deterministic subcriteria and diagnostic probes, so native counts are not the 77 report-case acceptance totals. Retry attempts are not double-counted. Live provider history: previous attempt HTTP 504; new single probe HTTP 503 after 7.324 seconds. Two cumulative calls out of eight authorised; no successful live response, token usage or cost measurement.')
for v in [('New xUnit extension','HistoricalCaseVerificationTests + isolated PG facts','4','2','2','Final cost-12 assertion fails; first TLS 349.529 ms fails; bootstrap/EF duplicate pass. TRX files preserved.'),('New Flutter extension','historical_cases_test.dart --machine','5','2','3','JSON model/navigation pass; rating, disabled-submit and jwt storage mechanisms conflict.'),('New Playwright extension','historical.playwright.config.ts','9','5','4','Final valid setup; real careers/auth plus scoped mocked UI. No-slot semantic retest remains failed.'),('New Python diagnostics','test_historical_acceptance.py; JUnit','5','3','2','Helpers pass only subcriteria. Completeness 75% versus 50%; raw retrieval-disclaimer diagnostic fails, not final-chat acceptance.')]:addrow(T[22],v)
cell(T[22].cell(0,0),'Suite / tool (baseline and new extensions)')
areaids=[('Backend / API',lambda i:i.startswith('TC-BK-')),('Database',lambda i:i.startswith('TC-DB-')),('React web',lambda i:i.startswith('TC-FE-')),('Flutter mobile',lambda i:i.startswith('TC-MOB-')),('Integration / E2E',lambda i:i.startswith('TC-E2E-')),('Component 1 (Jeyashankar)',lambda i:i.startswith('TC-C1')),('Component 2 (Buddhika)',lambda i:i.startswith('TC-C2-')),('Component 3 (Vithusan)',lambda i:i.startswith('TC-C3-')),('Component 4 (Wijayananda)',lambda i:i.startswith('TC-C4-'))]
row(T[23],0,['Area','Cases','VERIFIED/PASS','VERIFIED/FAIL','H/U or BLOCKED'])
for ri,(label,f) in enumerate(areaids,1):
 cc=[x for x in cases if f(x['test_id'])];n=collections.Counter(x['status'] for x in cc);row(T[23],ri,[label,len(cc),n['VERIFIED/PASS'],n['VERIFIED/FAIL'],f"{n['HISTORICAL/UNVERIFIED']} H/U; {n['BLOCKED']} BLOCKED"])
row(T[23],10,['Total','77','35','30','0 H/U; 12 BLOCKED'])
put(P[63],'Of the 49 previously historical cases, 39 now have scoped execution evidence: 12 matched and 27 differed from historical expectations. Ten are BLOCKED after investigation. Overall totals: 35 VERIFIED/PASS + 30 VERIFIED/FAIL + 12 BLOCKED = 77; zero HISTORICAL/UNVERIFIED and zero NOT EXECUTED. The earlier three query/headroom/cancellation expectation conflicts are retained. Failure counts include route/status/schema/precondition discrepancies awaiting requirements confirmation; they are not 30 confirmed defects.')
put(P[67],'Earlier baseline: xUnit 326 pass/2 skip, Node 19 pass, Flutter 49 pass, Playwright 111 pass/3 skip and Python 102 pass (607 passes). New extension runs: 23 native cases, 12 pass and 11 fail. Thirty-five report cases pass and 30 historical expectations fail; 12 remain blocked. The earlier spoof fix and alternate integrated Affidavit workflow retain their verified results. Original Deed/Approved workflow, full live AI, native-device acceptance and current branch CI remain incomplete.')
put(P[279],'Two authorised live Gemini requests, retries disabled: the earlier call returned HTTP 504 after 29.326 seconds; the new availability probe returned HTTP 503 after 7.324 seconds. Neither produced successful output, token usage or verified pricing; estimated cost is unavailable and the 1.5-second successful-response NFR remains BLOCKED. Cumulative calls are two of eight. No additional calls were made after the new provider failure. Prior mock-gateway timings remain historical mock measurements.')
# Do not invent a fix/closure for historical timezone defect.
for rr in T[25].rows[1:]:
 if rr.cells[0].text=='DEF-08':cell(rr.cells[-1],'New two-zone explicit business-time normalization passes; original implicit date-drift defect closure remains unverified. historical-resolution/evidence/timezone-conflict.json')
put(P[103],P[103].text+' Figures 6.31–6.36 add genuine historical-resolution output/UI captures. UI fixtures, live API pages and saved-log views are separately identified.')
figs=[]
native=json.loads((O/'evidence/native-capture.json').read_text());ui={x['test_id']:x for x in new};assigned=json.loads((O/'evidence/real-ui-assigned-steps.json').read_text());by={x.get('step'):x for x in assigned if 'step' in x}
items=[('password-cost-output.png','Historical BCrypt work-factor precondition fails',native['captured_utc'],'password-cost.log; password-cost.trx','Actual VSTest log in Chromium plain-text viewer; not a simulated terminal. Password hash length and verification pass, but cost 11 differs from historical 12. No product cryptographic rule changed.'),('TC-FE-05.png','React catalogue handles recorded HTTP 500',ui['TC-FE-05']['execution_date'],'playwright-historical.json','Actual React rendering with an explicitly mocked HTTP 500 response. The alert and Retry remain usable; this is UI error-handling evidence, not backend uptime.'),('TC-C2-06.png','React booking UI completes recorded approval fixture',ui['TC-C2-06']['execution_date'],'playwright-historical.json','Lawyer/date/slot selection and completion banner execute in real React. API approval is mocked; no PostgreSQL booking or live AI is inferred.'),('career-card.png','Senior Counsel career is visible from real PostgreSQL API',json.loads((O/'evidence/career-card-manifest.json').read_text())['captured_utc'],'historical-http.json; playwright-historical.json; career-card-manifest.json','Real API career creation returns 201 and the actual React careers page displays the vacancy. Synthetic local data; no API interception.'),('real-flutter-catalogue.png','Real Flutter catalogue reveals historical intake mismatch',by['real-flutter-catalogue']['at'],'real-ui-assigned-http.json; real-ui-assigned-steps.json','Real Flutter web login/catalogue executed against the isolated API. Current catalogue/service identity and sequential request/upload design differ from historical Deed intake. Two-PDF historical multipart creation was not completed.'),('real-react-clerk-review.png','Actual clerk review uses current statuses instead of Approved',by['real-react-clerk-review']['at'],'real-ui-assigned-http.json; real-ui-assigned-steps.json','Real API assignment and PostgreSQL attachment are visible in React. Status options exclude Approved and no Approve button exists; the historical action was not forced or replaced with COMPLETED.')]
for num,(im,title,date,src,ex) in enumerate(items,31):
 p=before(P[275],f'Figure 6.{num}. {title}');p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True;p.runs[0].bold=True
 p=before(P[275],'');p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(O/'evidence'/im),width=Inches(6.1))
 for txt,size in [('Execution/capture: '+date+' (UTC).',9),('Baseline: Wishwa2002/Intelligent-legal-service, se3110-qmse-testing, SHA 05359214df52; local uncommitted test overlay.',9),('Source: historical-resolution/evidence/'+src,8.5),(ex,10)]:
  p=before(P[275],txt,size);p.paragraph_format.keep_with_next=txt!=ex
 figs.append({'figure':f'6.{num}','title':title,'image':'historical-resolution/evidence/'+im,'source':['historical-resolution/evidence/'+s for s in src.split('; ')],'execution_date':date,'interpretation':ex})
addrow(T[42],['Figs 6.31–6.36','New historical acceptance replays','Native TRX/Flutter/pytest/Playwright; real API/PG/UI','historical-resolution/test-results/case-verification.json; evidence-index.json'])
before(P[276],'New TC-DB-01: one first EF TLS connection took 349.529 ms, exceeding its historical <200 ms connection target. This is not a new API p95 measurement and does not replace earlier controlled k6 results. No new load scan or coverage percentage is asserted.',10)
before(P[291],'Historical AI acceptance remains incomplete. The new configured Gemini availability call returned 503 after 7.324 seconds (previous call 504); cumulative count is two of eight. Date parsing, injection isolation, RRF comparison and actual hybrid retrieval executed only their documented subcriteria. Absence of a disclaimer in raw retrieved documents is not proof that final generated advice omits one. No expert-approved legal relevance dataset was available; full GC-01–GC-04 outcomes are unchanged.',10)
for ti in [19,44]:
 for rr in T[ti].rows[1:]:
  for c in rr.cells:
   if 'live-gemini.json' in c.text or 'provider' in c.text.lower():cell(c,c.text+' New provider probe: historical-resolution/evidence/live-gemini.json (503); no live acceptance PASS.')
row(T[45],0,['Member','Test cases','VERIFIED/PASS','VERIFIED/FAIL','H/U / BLOCKED','Component'])
memberparas=[(304,312,'C1'),(323,331,'C2'),(342,350,'C3'),(361,370,'C4')]
for ri,(pi,interpret,c) in enumerate(memberparas,1):
 ids=[x for x in cases if x['test_id'].startswith('TC-'+('C1' if c=='C1' else c+'-'))];n=collections.Counter(x['status'] for x in ids)
 for col,val in [(2,n['VERIFIED/PASS']),(3,n['VERIFIED/FAIL']),(4,f"0 H/U; {n['BLOCKED']} BLOCKED")]:cell(T[45].cell(ri,col),str(val))
 put(P[pi],f"{len(ids)} component-assigned cases: {n['VERIFIED/PASS']} VERIFIED/PASS, {n['VERIFIED/FAIL']} VERIFIED/FAIL, {n['BLOCKED']} BLOCKED, zero H/U. New replay work is AI-assisted; historical component assignments do not prove personal authorship.")
interpretations={312:'C1 new replays: fee is not persisted and historical availability toggle is absent. Live >90% recommendation and hiring rationale remain blocked. Query/headroom expectations are not changed. Jeyashankar must confirm fee/toggle requirements, scoring acceptance and workforce-rule precedence.',331:'C2 new replays: duplicate booking, real career visibility, scoped mocked booking UI and two browser timezone contexts pass. Historical immediate Confirmed and PUT cancellation expectations differ. Full live tomorrow-at-10 timestamp remains blocked. Buddhika must confirm status/API contract and original timezone scenario.',350:'C3 new replays expose historical service/checklist, upload/department/approval-route and completeness-ratio discrepancies. Current genuine 2 MiB upload succeeds, 14 MiB rejects; original spoof repair/regressions remain verified. Low-OCR live acceptance and atomic multi-file transaction remain blocked. Vithusan must confirm service mapping, bundle contents and transaction boundary.',370:'C4 new real React/API login/logout passes. Chat refusal, legal relevance/disclaimer and higher RRF precision remain blocked; helper results are not full provider acceptance. Wijayananda must confirm guardrail criteria and JWT/storage/redirect expectations with supporting requirements.'}
for i,s in interpretations.items():put(P[i],s)
for ri,c in enumerate(['C1','C2','C3','C4'],1):cell(T[66].cell(ri,2),T[66].cell(ri,2).text+' New historical replay outcomes: '+str(S['components'][c])+'. AI-assisted; assigned member review required.')
before(P[374],'Shared responsibility proposal (not effort attribution): Jeyashankar coordinates six catalogue/model cases; Buddhika seven scheduling/mobile-intake/status cases; Vithusan seven PostgreSQL/document-workflow cases; Wijayananda seven authentication/client-security cases. Full 27-ID allocation is in historical-resolution/audit/shared-responsibility-proposal.md. All four members must confirm the proposal and understand the test overlay.',10)
put(P[375],'Earlier QMSE work repaired and retested four defects and verified an alternate real Flutter/API/PostgreSQL/FastAPI/React Affidavit workflow. Those results remain intact. The historical-resolution phase investigated all 49 remaining cases: 39 scoped executions (12 pass, 27 fail) and ten blocked. Overall 77 cases: 35 VERIFIED/PASS, 30 VERIFIED/FAIL, 12 BLOCKED. The earlier 607 passing machine cases remain a dated baseline; new 23-case extensions have 12 pass and 11 fail. Many failures are historical contract/precondition discrepancies requiring approval, not confirmed product defects. Original Deed/Approved workflow, live AI, native-device acceptance, member confirmation and CI execution remain incomplete.')
put(P[378],P[378].text+' Also confirm API routes/statuses, cost-12 precondition, fee/card behaviour, checklist ratios and secure-storage/JWT/redirect claims before changing business rules or declaring product defects.')
put(P[382],P[382].text+' New provider probe is 503; remaining six calls are not automatically spent. New test overlay has not been published and requires separate publication authorisation.')
cell(T[67].cell(2,1),T[67].cell(2,1).text+' This revision adds AI-assisted 49-case investigation, native replays and requirement-discrepancy analysis; no new tests are credited to individual students.')
cell(T[67].cell(4,1),'Original spoof failure/retest, earlier 607-pass baseline, new 23-case native extensions (12 pass/11 fail), 77 report-case totals and provider 504/503 are evidence-linked. All 30 prior figures preserved; six genuine figures added. Manual member assessment remains pending.')
cell(T[68].cell(5,2),'77 original cases: 35 VERIFIED/PASS, 30 VERIFIED/FAIL, 12 BLOCKED; H/U and NOT EXECUTED zero. All 49 remaining historical cases investigated; blocked cases remain unfinished.')
cell(T[68].cell(9,2),'Earlier 102 deterministic/mocked tests retained. New five diagnostics: three pass/two fail, with subcriteria explicitly scoped. Live provider 504 then 503; no full live Golden Case verified.')
cell(T[68].cell(13,2),'VERIFIED native-count reconciliation: baseline 607 passes; new distinct native extensions 23 total, 12 pass, 11 fail. Report acceptance totals separately reconciled.')
cell(T[68].cell(14,2),T[68].cell(14,2).text+' New AI-assisted tests are supplied in historical-resolution/source-snapshot and remain uncommitted/unpublished.')
cell(T[68].cell(16,2),T[68].cell(16,2).text+' New isolated replay instructions: historical-resolution/audit/reproduction.md; exact environment and blockers recorded.')
before(P[408],'Historical-resolution overlay: clone and verify the published SHA above, then copy the hashed source-snapshot paths from this evidence package. These new tests are not yet on GitHub. Follow historical-resolution/audit/reproduction.md for fresh isolated database setup, opt-in variables, native commands and expected acceptance failures. Do not run seeding against production.',10)
put(P[287],P[287].text.replace('The 49 existing Flutter unit/widget cases also pass.','The earlier 49 Flutter unit/widget cases passed.'))
put(P[289],P[289].text.replace('All 102 final Python deterministic/mocked tests pass.','Earlier baseline: 102 Python deterministic/mocked tests passed.'))
# Prevent split rows and an empty explicit break from creating an orphan/blank page.
for table in D.tables:
 for rr in table.rows:
  pr=rr._tr.get_or_add_trPr()
  if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
for pp in list(D.paragraphs):
 nxt=pp._p.getnext()
 if not pp.text and pp._p.find('.//'+qn('w:br')) is not None and nxt is not None and ''.join(nxt.itertext()).startswith('Original group development repository:'):
  nxtpr=nxt.find(qn('w:pPr'))
  if nxtpr is None:nxtpr=OxmlElement('w:pPr');nxt.insert(0,nxtpr)
  nxtpr.append(OxmlElement('w:pageBreakBefore'));pp._p.getparent().remove(pp._p)
D.save(F/(name+'.docx'));(O/'audit/document-change-log.json').write_text(json.dumps(edits,indent=2));(O/'evidence/figure-index.json').write_text(json.dumps(figs,indent=2));print({'tables':len(D.tables),'images':len(D.inline_shapes),'edits':len(edits),'output':str(F/(name+'.docx'))})
