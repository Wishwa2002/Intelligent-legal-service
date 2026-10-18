from pathlib import Path
from datetime import datetime,timezone
import json,hashlib,shutil,zipfile,re,collections,xml.etree.ElementTree as ET
import pymupdf
from docx import Document
R=Path.cwd();O=R/'qmse-output';C=O/'completion-20261008';F=O/'final-report';name='SE3110_Testing_Report_GR42_Final';shutil.copy2(F/'render-completion'/f'{name}.pdf',F/f'{name}.pdf');old=Document(C/'audit/before.docx');D=Document(F/f'{name}.docx');pdf=pymupdf.open(F/f'{name}.pdf');checks={}
def check(k,b):
 checks[k]=bool(b)
 if not b:raise AssertionError(k)
check('docx_archive_opens',zipfile.ZipFile(F/f'{name}.docx').testzip() is None);check('tables_and_sections_preserved',len(D.tables)==len(old.tables)==71 and len(D.sections)==len(old.sections)==3);check('section_layout_preserved',[s._sectPr.xml for s in old.sections]==[s._sectPr.xml for s in D.sections]);check('20_genuine_images',len(D.inline_shapes)==20)
def media(p):
 with zipfile.ZipFile(p) as z:return {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
check('all_13_original_images_preserved',media(C/'audit/before.docx')<=media(F/f'{name}.docx'))
figs=json.load(open(O/'evidence/visual/figure-index.json'))+json.load(open(C/'evidence/figure-index.json'));placement=[]
for f in figs:
 title=f"Figure {f['figure']}. {f['title']}";found=[i+1 for i,p in enumerate(pdf) if title in p.get_text()];check('unique_caption_'+f['figure'],len(found)==1);p0=pdf[found[0]-1];check('image_caption_same_page_'+f['figure'],bool(p0.get_image_info()));check('image_file_exists_'+f['figure'],(O/f['image']).is_file())
 for src in f['source'].split('; '):
  if not src.startswith('https://'):check('figure_source_exists_'+src,(O/src).is_file())
 placement.append({'figure':f['figure'],'page':found[0]})
check('20_pdf_image_placements',sum(len(p.get_image_info()) for p in pdf)==20);check('images_inside_page',all(p.rect.contains(pymupdf.Rect(im['bbox'])) for p in pdf for im in p.get_image_info()));check('page_numbers_present',all(f'Page {i+1}' in p.get_text() for i,p in enumerate(pdf)))
for p,h in json.load(open(O/'audit/input-hashes.json')).items():check('original_unchanged_'+Path(p).name,hashlib.sha256(Path(p).read_bytes()).hexdigest()==h)
cs=json.load(open(C/'test-results/case-verification.json'));check('77_distinct_cases',len(cs)==len({c['test_id'] for c in cs})==77);check('case_totals',dict(collections.Counter(c['status'] for c in cs))=={'VERIFIED':23,'FAILED':3,'BLOCKED':2,'HISTORICAL/UNVERIFIED':49})
for c in cs:
 for e in c['evidence']:check('case_evidence_exists_'+e,(R/e).is_file())
check('all_57_historical_reviewed',len(json.load(open(C/'audit/historical-57-review.json')))==57)
ns={'t':'http://microsoft.com/schemas/VisualStudio/TeamTest/2010'};c=ET.parse(C/'evidence/backend-final-verified.trx').find('.//t:Counters',ns).attrib;check('xunit_328_326_0_2',c['total']=='328' and c['passed']=='326' and c['failed']=='0' and sum(x.attrib['outcome']=='NotExecuted' for x in ET.parse(C/'evidence/backend-final-verified.trx').findall('.//t:UnitTestResult',ns))==2);ps=json.load(open(C/'evidence/playwright-final.json'))['stats'];check('playwright_111_0_3',ps['expected']==111 and ps['unexpected']==0 and ps['skipped']==3);ai=ET.parse(C/'evidence/ai-completion.xml').find('testsuite');check('pytest_102_0',int(ai.attrib['tests'])==102 and int(ai.attrib['failures'])==0 and int(ai.attrib['errors'])==0)
check('integrated_pass',json.load(open(C/'evidence/integrated-result.json'))['status']=='PASS');check('live_only_one_attempt',json.load(open(C/'evidence/live-gemini.json'))['cumulative_attempted_calls']==1)
toc=json.load(open(C/'audit/toc-pages.json'));norm=lambda s:re.sub(r'\s+',' ',s).strip();check('54_toc_targets_match',len(toc)==54 and all(norm(x['title']) in norm(pdf[x['page']-1].get_text()) for x in toc.values()))
# Compare configured private secret values without printing them. Private files are outside package.
secrets=[]
for p in Path('/private/tmp/qmse-completion-control').glob('*.json'):
 if p.name in ['settings.json','performance.json']:
  v=json.load(open(p));secrets += [x for k,x in v.items() if k in ['password','key','token'] and isinstance(x,str) and len(x)>12]
for line in (R/'ai-service/.env').read_text().splitlines():
 if '=' in line:
  k,v=line.split('=',1)
  if any(s in k for s in ['API_KEY','SECRET']):v=v.strip().strip('"\'');secrets += [v] if len(v)>16 else []
hits=[];scanned=0
for p in O.rglob('*'):
 if not p.is_file() or p.suffix.lower() in ['.zip','.png','.pdf','.docx'] or any(x.startswith('render') or x in ['.dotnet','__pycache__'] for x in p.relative_to(O).parts):continue
 try:text=p.read_text();scanned+=1
 except (UnicodeError,OSError):continue
 if any(s in text for s in secrets):hits.append(str(p.relative_to(R)))
check('configured_secrets_not_in_text_evidence',not hits)
result={'checked_utc':datetime.now(timezone.utc).isoformat(),'all_checks_passed':all(checks.values()),'pdf_pages':len(pdf),'figures':20,'preserved_figures':13,'tables':71,'figure_pages':placement,'secret_text_files_scanned':scanned,'checks':checks,'visual_review':{'pages_inspected':list(range(1,len(pdf)+1)),'method':'All-page rendered contact sheets; changed pages inspected again after formatting corrections; original source images/outputs retained.','scope':'Screenshot readability, layout/table continuity, caption/image position, cover, TOC and footers. Source screenshot crops are intentional; full native reports retained.'}}
(C/'audit/final-quality-check.json').write_text(json.dumps(result,indent=2));print(json.dumps({k:v for k,v in result.items() if k not in ['checks','visual_review','figure_pages']},indent=2))
