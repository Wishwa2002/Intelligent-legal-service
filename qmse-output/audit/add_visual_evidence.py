from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
R=Path(__file__).resolve().parents[2];O=R/'qmse-output';D=Document(O/'final-report/previous/SE3110_Testing_Report_GR42_Final.docx')
# Source files and dates are execution metadata, not image-capture dates.
figs=[
('01-xunit-result','xUnit execution totals','2026-10-08 02:56:37 UTC','test-results/backend-final.log; evidence/backend/backend-final.trx','313 total: 310 passed, 1 failed, 2 skipped. Exit code 1. Skips include live-provider transport and the evidence-host harness.',6.26),
('02-spoofed-pdf-failure','DEF-12: spoofed PDF rejection fails','2026-10-08 02:56:37 UTC','test-results/backend-final.log; evidence/backend/backend-final.trx','Expected ArgumentException was not thrown for 32 inert ASCII bytes beginning MZ, supplied as sample.pdf/application/pdf. DEF-12 remains OPEN; no successful remediation retest exists.',6.26),
('03-node-result','React project: Node test output','2026-10-08 02:42:14 UTC','test-results/frontend-node.log; test-results/frontend-node.json','The actual npm run test:member1 execution passed all 19 subtests. These are Node routing/helper checks, not React Testing Library component tests.',6.26),
('04-flutter-result','Flutter machine runner: successful completion','2026-10-08 03:15:24 UTC','test-results/mobile-minimal.log; evidence/mobile/flutter-events.jsonl','Visible final widget cases succeed and the runner reports success=true. The full machine stream contains 49 visible unit/widget tests across six files; no device or integrated workflow is inferred.',6.26),
('05-playwright-report','Playwright HTML execution report','2026-10-08 02:42:40 UTC','evidence/e2e/playwright-report/index.html; evidence/e2e/playwright.json','109 passed and three skipped, across 14 specification files (112 cases). The report All tab excludes skipped cases. Most tests intercept API requests with mocks.',6.26),
('06-react-schedule','Actual React schedule editor under Playwright','2026-10-08 02:42:40 UTC (suite start)','evidence/e2e/schedule-1280.png; evidence/e2e/playwright.json','Focused browser view of the preserved application screenshot. Weekly hours, duration and synthetic 2030 leave data are visible. API responses were mocked: this is UI evidence, not the blocked Flutter-to-clerk E2E workflow.',5.35),
('07-zap-summary','OWASP ZAP passive scan summary','2026-10-08 02:54:12 UTC','evidence/security/zap-passive-report.html; evidence/security/scan-metadata.json','Six deliberate requests to isolated loopback HTTP; no active scan or spider. Two alert types: one Low and one Informational. Zero High/Medium alerts in this scope does not establish absence of vulnerabilities.',6.26),
('08-zap-alert-details','ZAP alert instances and missing-header detail','2026-10-08 02:54:12 UTC','evidence/security/zap-passive-report.html; evidence/security/zap-alerts.json','The Low missing X-Content-Type-Options alert has three instances; authentication discovery has one Informational instance. Fixture headers differ from full production startup. The report false-positive count is not a manual disposition.',6.26),
('09-k6-execution','k6 controlled local execution completes','2026-10-08 02:52:15 UTC','test-results/performance-final.log; test-results/performance-final.json','Actual output records completion of 20 iterations with two VUs over ten seconds. The fixture used EF InMemory and mocked AI; PostgreSQL and live Gemini performance were not measured.',6.26),
('10-k6-metrics','k6 original machine metrics','2026-10-08 02:52:15 UTC','evidence/performance/k6-summary.json; evidence/performance/k6-raw-final.json','HTTP failure rate is 0 across 81 requests. Appointment-list p50 is 1.263 ms, p95 23.507 ms, average 3.473 ms. The rate metric fails=81 means 81 false samples for http_req_failed, not 81 failed requests.',6.26),
('11-coverage','Coverlet results displayed by ReportGenerator','2026-10-08 02:56:37 UTC (test run)','evidence/backend/coverage-html/index.html; evidence/backend/e95ff97b-fff2-40aa-993d-8317a0ac8565/coverage.cobertura.xml','Line coverage is 23%. Coverlet XML root records 48.91% branches (1326/2711); ReportGenerator displays 48.7% (1293/2655) after its own aggregation. Both originals are retained; this rendering is not a new execution or coverage increase.',6.26),
('12-ai-tests','pytest deterministic and mocked AI suite','2026-10-08 02:47:35 UTC','test-results/ai-final.log; evidence/ai/pytest-final.xml','98 passed, with nine warnings. Synthetic provider initialization and mocked/deterministic tests do not demonstrate live Gemini accuracy, Golden Case success or response latency.',6.26),
('13-github-mobile-ci','GitHub Actions: historical Mobile CI success','2026-10-05 23:48:54 UTC (run created)','https://github.com/Jeyamdev/legalease-ai/actions/runs/37390618065; evidence/ci/github-actions-runs.json','The real GitHub page shows Success for base commit 8a9b8a2, duration 2m18s. It does not cover uncommitted local test additions; other recorded runs failed. Job logs require sign-in.',6.26),
]
anchor=next(p for p in D.paragraphs if p.text=='7. Non-Functional Test Analysis')
intro=next(p for p in D.paragraphs if p.text.startswith('Tool evidence is supplied'))
intro.runs[0].text='Tool evidence is supplied as original machine reports, logs and the following 13 genuine screenshots. Saved outputs are displayed in Chromium without changing their content; browser font, wrapping and viewport are adjusted for readability. These are artifact views, not recreated terminal windows. Execution dates appear below each figure; capture dates and source hashes are in evidence/visual/capture-manifest.json. All source paths are relative to qmse-output unless a URL is shown.'
for run in intro.runs[1:]:run.text=''
# Correct the evidence table figure references while retaining every original result and row.
refs=['Fig 6.1-6.2','Fig 6.3','Fig 6.4','Fig 6.5','Fig 6.11','Fig 6.9-6.10','Fig 6.7-6.8','Fig 6.12','Fig 6.13','Fig 6.6; E2E blocked']
for t in D.tables:
 if t.cell(0,0).text=='Ref':
  for row,ref in zip(t.rows[1:],refs):
   p=row.cells[0].paragraphs[0];p.runs[0].text=ref
   for run in p.runs[1:]:run.text=''
# Start deliberate figure groups on fresh pages; keep each caption with its image.
group_starts={1,3,4,5,6,7,8,9,11,12}
manifest=[]
for n,(id,title,date,source,explain,width) in enumerate(figs,1):
 p=anchor.insert_paragraph_before();p.paragraph_format.page_break_before=n in group_starts;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(5)
 r=p.add_run(f'Figure 6.{n}. {title}');r.bold=True;r.font.name='Calibri';r.font.size=Pt(11);r.font.color.rgb=RGBColor.from_string('17365D')
 p=anchor.insert_paragraph_before();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
 image=O/f'evidence/visual/{id}.png';p.add_run().add_picture(str(image),width=Inches(width))
 p=anchor.insert_paragraph_before();p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True;r=p.add_run('Execution: '+date+'.');r.font.name='Calibri';r.font.size=Pt(9)
 p=anchor.insert_paragraph_before();p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True;r=p.add_run('Source: '+source);r.font.name='Calibri';r.font.size=Pt(8.5)
 p=anchor.insert_paragraph_before();p.paragraph_format.space_after=Pt(9);p.paragraph_format.keep_with_next=False;r=p.add_run(explain);r.font.name='Calibri';r.font.size=Pt(10)
 manifest.append({'figure':f'6.{n}','title':title,'execution_date':date,'source':source,'image':str(image.relative_to(O)),'interpretation':explain})
p=anchor.insert_paragraph_before('Unavailable evidence: no successfully executed full Flutter → ASP.NET → PostgreSQL → AI → React workflow or live Golden Case run is present. No application-step images or successful spoofed-PDF retest have been manufactured. Existing blocker records and original logs remain authoritative.');p.paragraph_format.space_after=Pt(9)
for r in p.runs:r.font.size=Pt(10)
D.save(O/'final-report/SE3110_Testing_Report_GR42_Final.docx')
(O/'evidence/visual/figure-index.json').write_text(json.dumps(manifest,indent=2))
print('Added',len(D.inline_shapes),'figures; retained',len(D.tables),'tables')
