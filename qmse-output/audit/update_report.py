from pathlib import Path
from copy import deepcopy
import json,re,collections
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[2];O=R/'qmse-output';D=Document('/Users/jeyam/Downloads/SE3110_Testing_Report_GR42.docx'); blocks=list(D.element.body); S=json.load(open(O/'test-results/summary.json'));cases={x['test_id']:x for x in json.load(open(O/'test-results/case-verification.json'))};M=json.load(open(O/'evidence/performance/k6-summary.json'))['metrics'];snap=json.load(open(O/'audit/snapshot.json'));classcounts=json.load(open(O/'audit/backend-class-counts.json'))
def p(i):return Paragraph(blocks[i],D)
def t(i):return Table(blocks[i],D)
def setp(para,text):
 if para.runs:
  para.runs[0].text=text
  for run in para.runs[1:]:run.text=''
 else:para.add_run(text)
def setcell(cell,text):
 ps=cell.paragraphs
 setp(ps[0],str(text))
 for old in ps[1:]:old._element.getparent().remove(old._element)
def cell(i,r,c,text):setcell(t(i).cell(r,c),text)
def row(i,r,vals):
 for c,v in enumerate(vals):cell(i,r,c,v)
def addafter(anchor,text,style='Normal'):
 el=OxmlElement('w:p');anchor._element.addnext(el);para=Paragraph(el,D);para.style=style;para.add_run(text);return para
for table in D.tables:
 h=[c.text for c in table.rows[0].cells]
 if h[:2]==['Test ID','Feature / Type']:
  for r in table.rows[1:]:
   rec=cases[r.cells[0].text]; status=rec['status'];setcell(r.cells[6], 'H/U' if status=='HISTORICAL/UNVERIFIED' else status)
   if status=='HISTORICAL/UNVERIFIED':
    setcell(r.cells[5],'Historical result, UNVERIFIED. '+r.cells[5].text);setcell(r.cells[7],'audit/original-report-text.txt; historical claim only; no matching current execution artifact')
   else:
    setcell(r.cells[2],rec['environment'].replace('macOS ARM64 isolated fixtures; ','')+'; '+rec['preconditions']);setcell(r.cells[3],rec['exact_inputs']);setcell(r.cells[5],rec['observed']);setcell(r.cells[7],rec.get('automation_case','')+'; '+rec['evidence'][0].replace('qmse-output/','')+'; '+rec['limitations'])
# Minimal requested cleanup and factual replacements by original block identity.
replacements={
28:'Keep automated scripts and evidence in the repository so every test can be reproduced and explained.',
48:'Each case records design, expected outcome, actual observation and evidence. VERIFIED means a scoped fresh result supported by an execution artifact; FAILED means an observed defect or a report expectation conflict (identified explicitly). BLOCKED and NOT EXECUTED mean no valid execution result. H/U means HISTORICAL/UNVERIFIED: the original reported result is retained as history, not a current PASS. Current automation-case names and exact setups are in qmse-output/test-results/case-verification.json. Fixture dates such as 2030-01-14 are synthetic inputs, not execution dates.',
49:'3.1 Backend / API tests (xUnit, isolated Kestrel HTTP fixtures)',53:'3.3 React web tests (Node runner, Playwright)',
58:'Designed workflow: mobile intake, API mediation, persistence, AI checklist review, React clerk approval and mobile synchronization. BLOCKED: this complete workflow was not reproduced. Historical observations below have no matching integrated execution artifacts; UI mocks and deterministic agent-node tests do not establish this cross-platform workflow.',
80:'Current results come from saved TRX, pytest XML, Playwright JSON/HTML and Node output. The original SE3090 report was not supplied and its reported results cannot substitute for these executions. See qmse-output/test-results/test-execution-summary.md for commands, timestamps, environment, durations and exits. Golden Cases GC-01–04 remain historical/unverified; Python suite results are deterministic/mocked.',
84:'The 77 designed report cases comprise 14 scoped VERIFIED results, 4 FAILED expectations (three requirement conflicts and one spoofed-upload defect), 2 BLOCKED cases and 57 historical/unverified cases. These counts differ from machine-suite totals because multiple assertions/theory inputs may support one report case. Security and performance executions are described separately; no live Golden Case accuracy is verified.',
89:'The backend final suite records 310 passed, 1 failed and 2 skipped out of 313 cases, including a skipped scanner-host harness. The rejection test for spoofed PDF content fails. Node: 19 passed; Playwright: 109 passed and 3 skipped; Python: 98 deterministic/mocked cases passed. Flutter: 49 unit/widget tests passed; the complete integrated workflow remains blocked. See case-level scope and historical qualifications before assessing requirement coverage.',
91:'The original defect narratives and claimed fixes are retained as historical information. Commits were checked where cited, but changes do not prove successful retesting. Current findings, including the reproduced upload weakness and report expectation conflicts, are documented in qmse-output/audit/defect-verification.md. No application fix was made during this verification task.',
142:'Tool evidence is supplied as machine reports and logs in qmse-output. The table identifies actual artifacts and unresolved gaps. No placeholder screenshot or source-code reference is presented as a successful execution. Original report had no inline screenshots.',
146:'A controlled k6 run used two virtual users for 10 seconds against a loopback ASP.NET fixture with EF InMemory and mocked AI transport. It completed 20 iterations and 81 HTTP requests (including setup), with 0% HTTP failures and aggregate throughput 7.823 requests/second. Search, details and appointment-list fixture p95 values were below 200 ms. These measurements do not certify the hosted application or PostgreSQL performance.',
147:'The original login p95 of 180 ms has no measurement artifact. Login latency was not benchmarked in the new run. DEF-07 is a monitoring observation, not a demonstrated breach of the 200 ms budget. Hash cost and latency should be measured with the actual configured deployment.',
148:'The AI gateway measurement uses a mocked HTTP handler and must not be reported as Gemini latency. The original 1,450 ms live-AI p95 and 35 rps are historical/unverified. Live-AI acceptance against the 1.5-second target remains blocked pending approved fixtures, credentials and a bounded call budget.',
149:'The benchmark read appointment listings; it did not establish 280 rps for transactional booking. Isolated PostgreSQL concurrency tests verify one winning occupancy in their specific two-request scenario; they do not prove all load conditions.',
151:'Security checks used real JWT/controller fixtures and OWASP ZAP 2.17.0 passive inspection of six deliberate local requests. ZAP reported three Low missing-header alerts and one Informational authentication-discovery alert. The upload service accepts inert non-PDF bytes labelled as PDF: the rejection test fails. No active or production scan was performed; SQL-like input handling in EF InMemory is not proof of relational injection resistance.',
153:'Mocked timeout injection returns a generic 503 and subsequent AI health responds. An independent direct-booking test succeeds without an AI dependency. This supports scoped safe-failure/decoupling checks; no simultaneous complete-platform outage and recovery experiment was executed.',
155:'The browser suite executes actual React UI with mostly mocked API traffic, including responsive recurring-schedule views. Flutter: 49 unit/widget tests passed across six files after official SDK resource recovery. Complete web/mobile API equivalence, physical-device compatibility and accessibility conformance therefore remain unverified; Lighthouse/axe remain recommended.',
157:'The Python suite checks deterministic routing, retrieval, structured data, guardrails and human-decision nodes using mocks/synthetic fixtures. All 98 final tests pass. GC-01–04 live accuracy claims lack a reproducible labeled dataset, scoring denominator and matching run log, so no 98.9% mean accuracy or live-Gemini success rate is retained as verified.',
162:'The original component assignments and member disclosures are retained. Actual source/test files and all 20 cited commits were checked; personal authorship, effort and responsibilities still require member confirmation. This task’s AI-assisted executions cannot be attributed to an individual student without confirmation.',
164:'Section 3 status counts are report-case counts. Section 10 lists observed machine cases by named classes; these are not proof of personal ownership or effort percentages.',
172:'Isolated Kestrel HTTP fixtures exercise real controller routing, model binding and JWT authorization. Microsoft.AspNetCore.Mvc.Testing/WebApplicationFactory is not declared in the current test project.',
174:'Current recommendation eligibility and score validation use recorded data. Ranking points are recorded experience (0–70); the claimed weighted percentage formula is not implemented.',
177:'13 cases documented: 5 VERIFIED, 2 FAILED report expectation conflicts, 6 HISTORICAL/UNVERIFIED. The current-contract tests for the two expectation conflicts pass; approved boundaries/rules need confirmation.',
187:'Fresh component checks cover license duplication, Customer authorization, invalid specialization, current length/score bounds and mocked timeout handling. The report’s 10–1000 input boundary and below-20% hiring rule conflict with current code. Scores 96/92/91, weighted factors and GC-01 accuracy are historical/unverified. Related current user-link behavior is verified; original pre-fix defect claims require historical evidence.',
195:'Isolated PostgreSQL tests exercise current unique-slot/recurring overlap and concurrent mutation protection. The original claimed (LawyerId, AppointmentDate) constraint is not the current model.',
197:'The original 20-string live date-classification evaluation is historical/unverified. Deterministic schedule and conflict cases are traceable to current test artifacts.',
200:'12 cases documented: 3 VERIFIED, 1 FAILED report expectation conflict and 8 HISTORICAL/UNVERIFIED. The cancellation test records actual behaviour on both sides of 24 hours; current code imposes no notice cutoff.',
210:'Fresh service/database checks demonstrate one winning concurrent occupancy, outside-hours rejection (409), and direct booking without AI. All three cancellation offsets succeed, conflicting with the report’s below-24-hour rejection expectation. Original HTTP alternatives, two-zone UTC-fix retest and live scheduling Golden Case percentages are not independently verified.',
218:'The original zero-orphan rollback result is historical/unverified. A three-file atomic application filing/rollback operation was not identified and TC-C3-12 remains blocked.',
223:'12 cases documented: 3 VERIFIED, 1 FAILED (spoofed PDF acceptance), 1 BLOCKED rollback case and 7 HISTORICAL/UNVERIFIED.',
233:'Fresh file-size checks accept exactly 10 MiB and reject 10 MiB + 1 byte. Unknown-service creation is rejected and Customer status completion is denied with unchanged request state. Spoofed PDF content is accepted, contradicting DEF-12’s claimed fix. Seeder idempotence verifies 12 service records; this does not establish the live Document Agent or full clerk/mobile workflow.',
240:'Current xUnit tests exercise service-request logic and scoped JWT/authorization checks. AuthSecurityTests.cs is absent; do not assign its claimed ten cases to this member.',
241:'Fresh automated JWT checks use synthetic signing keys in an isolated HTTP fixture. The original jwt.io/Swagger tamper session is historical/unverified; no real token is sent to a third-party website.',
243:'GitHub Actions workflows are filtered to main/develop and component paths. Public metadata contains successes and failures; no universal every-push/green-pipeline claim is supported.',
246:'12 cases documented: 3 VERIFIED, 1 BLOCKED disclaimer case and 8 HISTORICAL/UNVERIFIED.',
256:'Fresh expired/invalid/tampered JWT and duplicate-signup tests pass, and Customer workflow approval is denied. Consultation disclaimer responses and original live prompt-injection/RAG accuracy claims remain unverified. Broad authentication/health remediation must not be inferred from a commit alone.',
259:'Observed class counts are derived from the final TRX. Original ownership assignments are historical statements pending member confirmation. New QMSE tests are AI-assisted work from this task, with no invented student attribution or percentage.',
261:'The verified results provide useful backend, relational, browser, deterministic AI, passive security and limited performance evidence. One backend test reproduces spoofed PDF acceptance. Three report expectations conflict with the current contract. The complete mobile-to-clerk workflow and live AI evaluation/latency remain blocked. Assessment completeness is therefore not established even though unsupported claims have been qualified and the evidence package is traceable.',
263:'Resolve the spoofed-file weakness against approved upload requirements, then rerun the retained failing rejection test; no business-logic fix was made here.',
264:'Confirm approved recommendation, workforce and cancellation rules before correcting the three report expectation conflicts. Repeat security scans on an isolated production-like setup; no public scan is authorized by this report.',
265:'Review the measured aggregate Coverlet coverage and add risk-focused assertions; do not impose an invented coverage threshold or equate coverage with correctness.',
266:'Retain the 49 passing Flutter unit/widget cases and execute one recorded cross-platform workflow using an isolated database and synthetic accounts.',
267:'Repeat the retained k6 script with representative PostgreSQL data and approved staging infrastructure. Measure live-AI latency separately with a bounded budget and approved inputs.',
268:'Attach a reproducible Golden Case dataset/scoring log, verify relevant CI runs for the tested commit, obtain member contribution confirmation, and add justified accessibility checks.',
280:'This checklist distinguishes verified document/evidence preparation from unresolved execution and assessment requirements. The legitimate individual demonstration/viva requirement is preserved; generic preparation instructions are replaced with actual status.',
285:'Reproduction requirements: .NET SDK 8.0.x, Node.js, a complete Flutter SDK, Python environment with ai-service requirements, and a dedicated local PostgreSQL cluster. Tested host is macOS ARM64; hosted Windows/Ubuntu results are historical except where CI metadata explicitly supports them.',
286:'Use synthetic local JWT/AI fixture configuration. Keep real database credentials and provider keys outside deliverables. Tests need no production data. Current pytest uses a synthetic initialization key and mocks, not live Gemini access.',
292:'Tests use synthetic example.test identities and test-only signing material. Production or confidential account credentials are excluded. A complete multi-role end-to-end fixture still requires isolated setup; the original SE3090 account-list reference is historical/unverified.'}
for i,text in replacements.items():setp(p(i),text)
# Source stack, actual tooling/environment, responsibilities.
row(16,1,['Backend','ASP.NET Core 8 Web API, EF Core 8, JWT HMAC-SHA256 and BCrypt; isolated tests do not certify production startup.'])
row(16,2,['Database','PostgreSQL via EF Core; current model exceeds 11 tables. Fresh relational tests use isolated local PostgreSQL 18. Cloud version/count and old unique-date claim are unverified.'])
row(16,4,['Mobile client','Flutter/Dart, ChangeNotifier, HTTP, shared_preferences, file_picker and flutter_dotenv. flutter_secure_storage is not declared in current pubspec.'])
row(16,5,['Agentic AI','Python/FastAPI/LangGraph/Pydantic and Gemini integrations. Tests use deterministic/mock inputs; original live model-version claims unverified.'])
row(16,6,['Pipeline','.github/workflows/backend-ci.yml, frontend-ci.yml and mobile-ci.yml; branch/path filters; observed CI run metadata in evidence/ci.'])
for r in t(33).rows[1:]:
 area=r.cells[0].text
 tool={'Backend / API':'xUnit, Moq, EF InMemory, isolated Kestrel HTTP hosts','Database':'xUnit + dedicated PostgreSQL/Npgsql/EF Core; no production access','React web':'Node test runner and Playwright; React static rendering and browser fixtures','Flutter mobile':'flutter_test: 49 unit/widget cases pass; no mocktail declared','Integration / E2E':'Scoped HTTP/PostgreSQL and browser-mock workflows verified; complete cross-platform workflow blocked','Performance':'k6 2.3.0, 2 VUs/10 s, local InMemory/mock fixture','Security':'xUnit HTTP/service probes; OWASP ZAP 2.17.0 passive local scan','Agentic AI':'pytest/Pydantic/LangGraph deterministic and mocked tests; live Golden Cases unverified'}.get(area)
 if tool:setcell(r.cells[2],tool)
row(36,1,['Operating system','Actual verification: macOS 26.4.1 ARM64. Original developer Windows/CI Ubuntu context is historical.'])
row(36,2,['Runtimes','.NET SDK 8.0.423; Node 22.23.2; Python 3.12 (AI venv). Flutter 3.47.6 official sources/resource recovery with matching cached test engine: unit/widget runtime verified.'])
row(36,3,['Services under test','Isolated Kestrel fixtures on random loopback ports; Playwright Vite 127.0.0.1:5175; ZAP proxy 55443. No public deployment tested.'])
row(36,4,['Database','Dedicated temporary PostgreSQL 18 cluster, 127.0.0.1:55442. Performance/security fixture uses EF InMemory. No Neon/production reset, seed or migration.'])
row(36,5,['Test accounts','Synthetic example.test Admin/Customer/Lawyer/Clerk fixtures; confidential credentials excluded.'])
row(36,6,['Version control',snap['branch']+' at '+snap['commit'][:12]+'. Requested remote report branch returned 404; no branch change, push or deployment.'])
for r in t(42).rows[1:]:setcell(r.cells[2],'Historical assigned responsibility; current files/counts checked in contribution audit. Personal ownership requires member confirmation.')
addafter(p(44),'The dated schedule below is the original plan, not evidence that activities occurred on those dates. Fresh execution times are recorded in qmse-output/test-results/*.json.')
# Golden Case accuracy qualification.
for r in t(69).rows[1:]:setcell(r.cells[4],'Not measured; historical '+r.cells[4].text+' unverified');setcell(r.cells[5],'H/U');setcell(r.cells[6],'No matching dataset/scoring/run artifact; audit/report-audit.md')
# Security current mappings.
row(72,1,['TC-SEC-01','JWT tampering','Fresh synthetic signature tamper on protected lawyer POST','401','401 in actual HTTP fixture','VERIFIED','Qmse_TC_C4_09...; evidence/backend/backend-final.trx'])
row(72,2,['TC-SEC-02','RBAC enforcement','Original clerk-list scenario not independently executed; related fresh staff/approval denials pass','403','Historical clerk-list claim unverified; related scoped checks verified','H/U','Fresh scope: Qmse_TC_C3_11 / Qmse_TC_C4_11'])
row(72,3,['TC-SEC-03','Password storage','Actual PasswordService verifies synthetic correct/wrong inputs','Correct verifies; wrong rejected; no plaintext','Scoped hashing checks pass; explicit 12-round expectation not configured','VERIFIED','Qmse_PasswordHashRejectsWrongPasswordAndReportsCost; backend TRX'])
row(72,4,['TC-SEC-04','Prompt-injection defence','Original live adversarial consultation','Safe refusal','No matching live refusal run; deterministic tests are narrower','H/U','evidence/ai/pytest-final.xml; audit limits'])
row(72,5,['TC-SEC-05','OWASP ZAP passive scan','6 requests; isolated loopback target; no active scan','Triage observed alerts','3 Low missing-header instances; 1 Informational auth identification','VERIFIED','evidence/security/zap-passive-report.html; scan-metadata.json'])
row(72,6,['TC-SEC-06','SQL-like input handling','Search: '+"' OR 1=1 --"+'; isolated EF InMemory','Safe handling; relational resistance requires actual SQL test','200, empty result; relational injection assurance unverified','VERIFIED (input only)','evidence/security/scan-metadata.json; Qmse_SqlLikeSearch...'])
row(72,7,['TC-SEC-07','Upload restriction','Exact 10 MiB/+1-byte probes; inert spoofed PDF','Size enforced; malicious type/content rejected','Size tests pass; spoof rejection fails; original full MIME scenario not executed','FAILED (spoof)','QmseDocumentTests; evidence/backend/backend-final.trx'])
# Performance table: discard unsupported historic metrics, add accurately scoped measured cases.
for r in t(75).rows[1:]:
 for c in [3,4,5]:setcell(r.cells[c],'Not measured')
 setcell(r.cells[6],'NOT EXECUTED');setcell(r.cells[7],'Original numbers unverified; no matching artifact')
def metric(name):
 v=M[name]['values'];return [f"{v['p(50)']:.3f} ms",f"{v['p(95)']:.3f} ms"]
row(75,1,['PERF-01','GET /api/lawyers/search (fixture)','< 200 ms',*metric('lawyer_search_ms'),'Aggregate 7.823','VERIFIED (fixture)','k6-summary.json; InMemory loopback only'])
row(75,5,['PERF-05','POST /api/lawyer-recommendations (mock gateway)','Live AI <1,500 ms',*metric('mock_ai_gateway_ms'),'Aggregate 7.823','Live BLOCKED; mock measured','No live Gemini; k6-summary.json'])
for vals in [['PERF-LOCAL-02','GET /api/lawyers/{id} (fixture)','<200 ms',*metric('lawyer_detail_ms'),'Aggregate 7.823','VERIFIED (fixture)','evidence/performance/k6-summary.json'],['PERF-LOCAL-03','GET /api/appointments (fixture)','<200 ms',*metric('appointment_list_ms'),'Aggregate 7.823','VERIFIED (fixture)','evidence/performance/k6-summary.json']]:
 r=t(75).add_row()
 for c,v in zip(r.cells,vals):setcell(c,v)
# Execution summary.
row(79,1,['Backend xUnit (.NET 8)','dotnet test with TRX + Coverlet',313,310,1,'2 skipped: live AI + evidence-host harness; 21.636 s wall, exit 1; isolated PostgreSQL and InMemory'])
row(79,2,['React/Node tests','npm run test:member1',19,19,0,'1.170 s wall; exit 0; static render/routing/HTTP fixture checks'])
row(79,3,['Flutter (flutter_test)','test --no-pub --machine','49','49','0','VERIFIED: 0 skipped; six files; 9.115 s event duration; 145.928 s runner wall; exit 0'])
row(79,4,['Playwright','npm test; JSON/HTML reports',112,109,0,'3 live-backend skips; 14 spec files; 68.243 s tool duration; mostly mocked API'])
row(79,5,['Python deterministic/mocked','pytest tests --junitxml',98,98,0,'1.833 s tool duration; synthetic initialization key; no live Gemini; GC percentages unverified'])
row(82,0,['Area','Cases','VERIFIED','FAILED','H/U or BLOCKED'])
counts=[[7,0,0,'7 H/U'],[4,0,0,'4 H/U'],[5,0,0,'5 H/U'],[5,0,0,'5 H/U'],[7,0,0,'7 H/U'],[13,5,2,'6 H/U'],[12,3,1,'8 H/U'],[12,3,1,'7 H/U + 1 BLOCKED'],[12,3,0,'8 H/U + 1 BLOCKED'],[77,14,4,'57 H/U + 2 BLOCKED']]
for n,values in enumerate(counts,1):
 for c,v in enumerate(values,1):cell(82,n,c,v)
for r in t(86).rows[1:]:setcell(r.cells[3],'Partial/UNVERIFIED requirement acceptance; see scoped case statuses')
# Historical defects: qualify every original FIXED/PASS and correct DEF-07/DEF-12.
for i in [95,98,101,104,107,110,113,116,119,122,125,128,131,134,137]:
 tab=t(i)
 for r in tab.rows:
  if r.cells[0].text=='Evidence and retest':setcell(r.cells[1],'Historical claim (UNVERIFIED): '+r.cells[1].text+'. Current scope: audit/defect-verification.md.')
  if r.cells[0].text=='Status':setcell(r.cells[1],'Historical reported status; current original retest UNVERIFIED. See defect verification.')
# Original defect summary must agree with detailed retest qualifications.
for rr in t(92).rows[1:]:
    defect=rr.cells[0].text
    setcell(rr.cells[4], 'OPEN' if defect=='DEF-12' else 'OBSERVATION' if defect=='DEF-07' else 'Historical claim / unverified')
    setcell(rr.cells[5], 'FAILED' if defect=='DEF-12' else 'UNVERIFIED')
# Original narrative row fixes remain clearly historical.
for i in [95,98,101,104,107,110,116,119,122,125,128,131,134,137]:
 for r in t(i).rows:
  if r.cells[0].text=='Fix / action':setcell(r.cells[1],'Historical reported action (not proof of fix/retest): '+r.cells[1].text)
for r in t(113).rows:
 if r.cells[0].text=='Status':setcell(r.cells[1],'OBSERVATION; historic login metric UNVERIFIED; not a demonstrated NFR breach.')
for r in t(128).rows:
 if r.cells[0].text=='Evidence and retest':setcell(r.cells[1],'Fresh service probe accepts inert MZ bytes labelled sample.pdf/application/pdf; rejection assertion fails. Qmse_DEF12_SpoofedPdfMustBeRejected; evidence/backend/backend-final.trx.')
 if r.cells[0].text=='Status':setcell(r.cells[1],'OPEN; retest FAILED. Claimed magic-byte fix contradicted by current code and execution.')
for i,detail in [(98,'Related user-link creation assertion passes; historical pre-fix reproduction unverified.'),(119,'Current non-admin career writes deny 401/403 in real HTTP fixture.'),(131,'Fresh seeder called twice: 12 services both times; live agent workflow not retested.')]:
 for r in t(i).rows:
  if r.cells[0].text=='Evidence and retest':setcell(r.cells[1],detail+' See audit/defect-verification.md and backend-final.trx.')
for r in t(140).rows[1:]:
 setcell(r.cells[1],r.cells[1].text+' historical');setcell(r.cells[2],'UNVERIFIED');setcell(r.cells[3],'DEF-12 reproduced; full historic status unverified' if r.cells[0].text in ['High','Total'] else 'Not adjudicated');setcell(r.cells[4],'Original retests UNVERIFIED; scoped checks in audit')
# Evidence rows use genuine paths or explicit blockers.
evidence=[('xUnit: 313 total, 310 pass, 1 fail, 2 skip','TRX + Coverlet','evidence/backend/backend-final.trx'),('React Node: 19 pass','Node runner','test-results/frontend-node.log'),('Flutter: 49 pass / 6 files','Flutter machine events','test-results/mobile-minimal.log; test-results/mobile-minimal.json'),('Playwright: 109 pass, 3 skip, 14 spec files','JSON + HTML','evidence/e2e/playwright.json; evidence/e2e/playwright-report/index.html'),('Coverage: 23.00% lines / 48.91% branches','Coverlet',S['backend']['coverage_file'].replace('qmse-output/','')),('Performance: local InMemory/mock','k6 2.3.0','evidence/performance/k6-summary.json; performance-report.md'),('Passive security: 4 alert instances','ZAP 2.17.0','evidence/security/zap-passive-report.html; security-report.md'),('Live Golden Cases: H/U; deterministic pytest 98 pass','pytest XML','evidence/ai/pytest-final.xml; original GC dataset/log absent'),('CI metadata: observed successes and failures','Public GitHub API','evidence/ci/github-actions-runs.json; current local additions not covered'),('Complete workflow: BLOCKED. Current recurring UI images use mocked API','Playwright screenshot','evidence/e2e/schedule-1280.png; unsupported-progress.png (not mobile-clerk workflow)')]
for n,(desc,tool,path) in enumerate(evidence,1):
 cell(143,n,1,desc);cell(143,n,2,tool);cell(143,n,3,path)
# AI control/metric tables remove unsupported generic assertions.
for r in t(158).rows[1:]:
 setcell(r.cells[2],'Scoped deterministic/mock evidence or blocker in case-verification.json; original whole-control claim unverified')
row(158,1,['Structured output','Pydantic contracts and current recorded-experience ranking','Scoped schema/ranking tests verified; no live schema-valid response percentage'])
row(158,5,['Failure recovery','Mock timeout → generic 503; health returns 200','TC-C113 verified mocked scope; claimed universal 10 s/back-off policy not established'])
for r in t(160).rows[1:]:setcell(r.cells[2],'NOT MEASURED');setcell(r.cells[3],'Original percentages historical/unverified; no labeled live evaluation dataset/run')
row(163,0,['Member','Test cases','VERIFIED','FAILED','H/U / BLOCKED','Component'])
for n,vals in enumerate([[13,5,2,'6 H/U'],[12,3,1,'8 H/U'],[12,3,1,'7 H/U + 1 BLOCKED'],[12,3,0,'8 H/U + 1 BLOCKED']],1):
 for c,v in enumerate(vals,1):cell(163,n,c,v)
# Preserve disclosures and assignments, qualify responsibility/old execution claims.
for idx in [169,192,215,238]:
 for r in t(idx).rows[1:]:setcell(r.cells[1],'Historical member-assigned claim; ownership confirmation pending: '+r.cells[1].text.replace('all passing','original pass claim unverified'))
for idx in [181,204,227,250]:
 for r in t(idx).rows[1:]:setcell(r.cells[3],'Historical AI-use disclosure; verification outcome UNVERIFIED: '+r.cells[3].text)
member_classes=[['FrontDeskMatchingTests','WorkforceHiringTests'],['AppointmentServiceTests','CareerServiceTests'],['DocumentationRequestServiceTests','ClerkServiceTests'],['ServiceRequestServiceTests']]
for n,names in enumerate(member_classes,1):
 counts=[sum(classcounts.get('LegalService.Tests.'+name,{}).values()) for name in names];cell(258,n,1,', '.join(f'{name}: {count} observed cases' for name,count in zip(names,counts)));cell(258,n,2,'Historical component assignment; personal testing work needs confirmation'+('; AuthSecurityTests.cs absent' if n==4 else ''));cell(258,n,3,str(sum(counts))+' observed; ownership unverified')
# CLEAR declaration preserves earlier disclosure while declaring this task accurately.
row(270,2,['Limits','Original report copied earlier figures and did not execute tests. This Codex-assisted revision ran local suites and recorded actual artifacts; mocked AI results, blockers and historical claims are identified. No production access, application fix, push, deployment or submission.'])
row(270,4,['Assessment','Claims were checked against the current repository and saved execution artifacts. A spoofed-upload assertion fails; report-contract conflicts and missing integrated/live evidence are disclosed. Final render QA is recorded separately.'])
row(270,5,['Reflection','Every reported result must be reproducible and understood by the responsible member. AI assistance does not establish personal ownership or replace checking the actual system.'])
addafter(p(269),'Revision assistance: Codex generated scoped test additions, ran tools, assembled evidence and edited this report on the actual logged execution date. Members must inspect, understand and confirm this work before submission.')
# Appendix compliance replaces preparation placeholders with observed status.
statuses=['VERIFIED: same repository/system; isolated tests','BLOCKED: complete cross-platform workflow not reproduced','Documented; approved SE3090 mapping still needs confirmation','VERIFIED document structure; actual environment corrected','Designed; 14 verified/4 failed/2 blocked/57 historical in 77 cases','Partially executed; full integration/live AI gaps disclosed','Partial: passive local security and fixture k6; live/deployment acceptance blocked','Scoped recovery/responsive UI and mobile widget tests; accessibility conformance unverified','Deterministic/mock suite verified; live Golden Cases H/U','Fresh scoped observations; original PASS claims qualified','DEF-12 reproduced; historical fixes/retests unverified','VERIFIED paths in evidence index; missing original evidence disclosed','VERIFIED machine summaries and report-case counts','Local test sources/scripts retained; commit/push not performed','20 cited commits exist; personal ownership requires confirmation','Replay files retained; live environment/full workflow blocked','NOT VERIFIED: individual demonstration/member understanding pending','VERIFIED updated declaration; member historical disclosures preserved']
for n,status in enumerate(statuses,1):cell(281,n,2,status)
# Reproduction commands use actual project/route; public scans removed.
for tab in D.tables:
 for r in tab.rows:
  for c in r.cells:
   old=c.text
   if 'dotnet build backend/LegalService.sln' in old:setcell(c,old.replace('dotnet build backend/LegalService.sln','dotnet build LegalService.sln'))
   if 'AuthSecurityTests' in old and 'dotnet test' in old:setcell(c,'AuthSecurityTests.cs absent. Use dotnet test backend/LegalService.Tests/LegalService.Tests.csproj --filter FullyQualifiedName~Qmse_TC_C4')
row(290,5,['Security scan','Run retained evidence/security/passive_scan.py only with an isolated loopback fixture and local ZAP proxy; no production target.'])
row(290,6,['Load test','Run retained evidence/performance/local-smoke.k6.js with a local Qmse_LocalEvidenceHost; mocked AI only.'])
# Remove all yellow placeholders without claiming successful execution.
for para in D.element.findall('.//'+qn('w:p')):
 for txt in para.findall('.//'+qn('w:t')):
  if txt.text and '[INSERT' in txt.text:txt.text=re.sub(r'\[INSERT[^\]]*\]','BLOCKED: matching execution artifact not available',txt.text)
 for hi in para.findall('.//'+qn('w:highlight')):
  hi.getparent().remove(hi)
# Repeat table headers; retain original page geometry and hierarchy, avoid fixed-height clipping.
for table in D.tables:
 trPr=table.rows[0]._tr.get_or_add_trPr()
 if trPr.find(qn('w:tblHeader')) is None:trPr.append(OxmlElement('w:tblHeader'))
 for r in table.rows:
  for el in list(r._tr.findall('.//'+qn('w:trHeight'))):el.getparent().remove(el)
settings=D.settings.element;up=settings.find(qn('w:updateFields'))
if up is None:up=OxmlElement('w:updateFields');settings.append(up)
up.set(qn('w:val'),'true')
# Keep cached TOC titles consistent with corrected headings.
for node in D.element.findall('.//'+qn('w:t')):
 if node.text:
  node.text=node.text.replace('3.1 Backend / API tests (xUnit, WebApplicationFactory)','3.1 Backend / API tests (xUnit, isolated Kestrel HTTP fixtures)').replace('3.3 React web tests (Vitest/Node runner, Playwright)','3.3 React web tests (Node runner, Playwright)')
# Verification note is near the report purpose, without a new section.
addafter(p(14),'Verification revision: 8 October 2026 (Asia/Colombo). Tested base '+snap['commit'][:12]+' on '+snap['branch']+'. Scope and blockers are explicit. Companion evidence paths are relative to qmse-output; original Word file remains unchanged.')
# Compact the existing TOC, preserve hyperlink anchors and native PAGEREF fields.
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
sdt=D.element.find('w:body/w:sdt',ns)
if sdt is not None:
 for el in sdt.findall('.//w:p',ns):
  para=Paragraph(el,D);para.paragraph_format.space_before=Pt(0);para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=Pt(10);para.paragraph_format.keep_with_next=False
  for rr in el.findall('.//w:r',ns):
   rp=rr.find(qn('w:rPr'))
   if rp is None:rp=OxmlElement('w:rPr');rr.insert(0,rp)
   sz=rp.find(qn('w:sz'))
   if sz is None:sz=OxmlElement('w:sz');rp.append(sz)
   sz.set(qn('w:val'),'18')
# Avoid orphaned small tables and partial reference lists.
for idx in [20,75,163,179,202,225,248]:
 tab=t(idx)
 for n,rr in enumerate(tab.rows):
  for cc in rr.cells:
   for para in cc.paragraphs:
    para.paragraph_format.keep_with_next=n<len(tab.rows)-1
    para.paragraph_format.space_after=Pt(0)
    if idx==75:
     for run in para.runs:run.font.size=Pt(8)
for idx in [271,272,273,274,275,276]:p(idx).paragraph_format.keep_with_next=True
# Correct inherited PASS/PLANNED colour cues to current status.
for tab in D.tables:
 headers=[cc.text for cc in tab.rows[0].cells]
 if 'Status' not in headers:continue
 col=headers.index('Status')
 for rr in tab.rows[1:]:
  cc=rr.cells[col];value=cc.text
  fill='FDE8E7' if 'FAILED' in value or value=='OPEN' else 'E6F1EB' if value.startswith('VERIFIED') else 'FFF3D8' if 'BLOCKED' in value or 'NOT EXECUTED' in value else 'EEF0F3'
  pr=cc._tc.get_or_add_tcPr();shd=pr.find(qn('w:shd'))
  if shd is None:shd=OxmlElement('w:shd');pr.append(shd)
  shd.set(qn('w:fill'),fill)
  ink='A12622' if 'FAILED' in value or value=='OPEN' else '1B6B3A' if value.startswith('VERIFIED') else '775A18' if 'BLOCKED' in value or 'NOT EXECUTED' in value else '495566'
  for para in cc.paragraphs:
   for run in para.runs:run.font.color.rgb=RGBColor.from_string(ink)
setp(p(14),'This report records tool-based testing and quality evaluation of the LegalEase AI system developed for SE3090. It preserves the SE3110 test plan, case designs, defects and contribution disclosures while distinguishing actual execution evidence from historical claims and unresolved requirements.')
for i in [182,205,228,251]:setp(p(i),'Historical defect and fix assignments')
# Correct misleading or unsafe duplicate reproduction rows.
row(290,5,['Flutter tests','cd mobile && flutter pub get && flutter test. Verified machine events: test-results/mobile-minimal.log; earlier setup attempts retained.'])
row(290,6,['Security scan','Use retained evidence/security/passive_scan.py with an isolated loopback fixture and local passive ZAP proxy.'])
row(290,7,['Load test','Use retained evidence/performance/local-smoke.k6.js with the local Qmse_LocalEvidenceHost; mock AI only.'])
# Final review corrections retained in the generator.
for tab in D.tables:
 for rr in tab.rows:
  if rr.cells[0].text=='Upload test (manual tool)':setcell(rr.cells[1],'Isolated fixture only: POST multipart /api/documentation-requests/{requestId}/files with a synthetic PDF exceeding 10 MiB; rejection expected. See QmseDocumentTests for the executed service probes.')
  for cc in rr.cells:
   if cc.text.startswith('Partial/UNVERIFIED requirement'):
    shd=cc._tc.get_or_add_tcPr().find(qn('w:shd'))
    if shd is None:shd=OxmlElement('w:shd');cc._tc.get_or_add_tcPr().append(shd)
    shd.set(qn('w:fill'),'EEF0F3')
    for para in cc.paragraphs:
     for run in para.runs:run.font.color.rgb=RGBColor.from_string('495566')
D.save(O/'final-report/SE3110_Testing_Report_GR42_Final.docx')
print('Saved final copy; original structure, cover, identities, references and legitimate viva requirements retained.')
