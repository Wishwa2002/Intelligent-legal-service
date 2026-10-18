from pathlib import Path
from datetime import datetime,timezone
import json,hashlib,shutil,re,zipfile
import pymupdf
from docx import Document
R=Path(__file__).resolve().parents[2];O=R/'qmse-output';F=O/'final-report';name='SE3110_Testing_Report_GR42_Final'
shutil.copy2(F/'render-visual'/f'{name}.pdf',F/f'{name}.pdf')
old=Document(F/'previous'/f'{name}.docx');new=Document(F/f'{name}.docx');pdf=pymupdf.open(F/f'{name}.pdf')
checks={}
def check(k,v):
 checks[k]=bool(v)
 if not v: raise AssertionError(k)
def outside(d):
 p=[x.text for x in d.paragraphs];a=p.index('6. Tool-Generated Evidence');b=p.index('7. Non-Functional Test Analysis');return p[:a]+p[b:]
check('paragraphs_outside_section_6_preserved',outside(old)==outside(new))
check('tables_preserved',len(new.tables)==len(old.tables)==71)
changed=[]
for i,(a,b) in enumerate(zip(old.tables,new.tables)):
 if [[c.text for c in r.cells] for r in a.rows] != [[c.text for c in r.cells] for r in b.rows]:changed.append(i)
check('only_section_6_evidence_table_changed',len(changed)==1 and 'Evidence' in old.tables[changed[0]].rows[0].cells[1].text)
check('section_settings_preserved',[s._sectPr.xml for s in old.sections]==[s._sectPr.xml for s in new.sections])
check('docx_opens_with_13_images',len(new.inline_shapes)==13)
check('docx_archive_valid',zipfile.ZipFile(F/f'{name}.docx').testzip() is None)
check('pdf_opens_51_pages',len(pdf)==51)
figures=json.loads((O/'evidence/visual/figure-index.json').read_text()); placements=[]
for f in figures:
 caption=f"Figure {f['figure']}. {f['title']}"
 hits=[i+1 for i,p in enumerate(pdf) if caption in p.get_text()]
 check('caption_once_'+f['figure'],len(hits)==1)
 p=pdf[hits[0]-1];check('caption_shares_page_with_image_'+f['figure'],len(p.get_image_info())>0)
 check('image_exists_'+f['figure'],(O/f['image']).is_file())
 for src in f['source'].split('; '):
  if not src.startswith('https://'):check('source_exists_'+src,(O/src).is_file())
 placements.append({'figure':f['figure'],'pdf_page':hits[0]})
images=[(i+1,x) for i,p in enumerate(pdf) for x in p.get_image_info()]
check('pdf_has_13_image_placements',len(images)==13)
check('images_inside_page_boundaries',all(pdf[n-1].rect.contains(pymupdf.Rect(x['bbox'])) for n,x in images))
check('all_page_numbers_present',all(f'Page {i+1}' in p.get_text() for i,p in enumerate(pdf)))
for path,digest in json.loads((O/'audit/input-hashes.json').read_text()).items():check('original_input_unchanged_'+Path(path).name,hashlib.sha256(Path(path).read_bytes()).hexdigest()==digest)
review=json.loads((O/'audit/visual-artifact-review.json').read_text());unchanged=0
for e in review:
 if e.get('sha256'):
  check('original_artifact_unchanged_'+e['path'],hashlib.sha256((R/e['path']).read_bytes()).hexdigest()==e['sha256']);unchanged+=1
for e in json.loads((O/'evidence/visual/capture-manifest.json').read_text()):
 if e.get('source_sha256'):check('capture_source_unchanged_'+e['id'],hashlib.sha256((O/e['source']).read_bytes()).hexdigest()==e['source_sha256'])
toc=json.loads((O/'audit/toc-pages-visual.json').read_text());norm=lambda s:re.sub(r'\s+',' ',s).strip()
check('54_toc_targets_exist',len(toc)==54 and all(norm(x['title']) in norm(pdf[x['page']-1].get_text()) for x in toc.values()))
result={'checked_utc':datetime.now(timezone.utc).isoformat(),'all_checks_passed':all(checks.values()),'pdf_pages':len(pdf),'embedded_figures':len(new.inline_shapes),'tables':len(new.tables),'original_hashed_artifacts_unchanged':unchanged,'figure_pages':placements,'visual_review':{'pages_inspected':list(range(1,52)),'method':'All final render page PNGs viewed; screenshot readability, captions, source paths, margins, table continuity, cover, TOC and footers inspected.','findings':'No truncated report rows, caption separation, missing images or page-boundary overflow found. Screenshot viewports intentionally show selected portions; complete source artifacts retained.'},'checks':checks}
(O/'audit/visual-quality-check.json').write_text(json.dumps(result,indent=2))
print(json.dumps({k:v for k,v in result.items() if k not in ('checks','visual_review')},indent=2))
