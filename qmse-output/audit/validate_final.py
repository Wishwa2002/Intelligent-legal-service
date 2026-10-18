from pathlib import Path
import json,re,hashlib,zipfile,shutil,datetime
from docx import Document
import pymupdf
R=Path(__file__).resolve().parents[2];O=R/'qmse-output';F=O/'final-report';src=F/'render-final/SE3110_Testing_Report_GR42_Final.pdf';dest=F/src.name;shutil.copy2(src,dest)
d=Document(F/'SE3110_Testing_Report_GR42_Final.docx');original=Document('/Users/jeyam/Downloads/SE3110_Testing_Report_GR42.docx');pdf=pymupdf.open(dest);texts=[re.sub(r'\s+',' ',p.get_text()).strip() for p in pdf];body=' '.join(texts)
cases=json.load(open(O/'test-results/case-verification.json'));assert len(cases)==77
for c in cases:
 for ev in c['evidence']:assert (R/ev).is_file(),ev
 assert c['test_id'] in body,c['test_id']
toc=json.load(open(O/'audit/toc-pages.json'))
for target in toc.values():
 hits=[i+1 for i,p in enumerate(texts) if i>=2 and target['title'] in p];assert hits and hits[0]==target['page'],target
assert len(d.tables)==len(original.tables)
assert not re.search(r'\[INSERT[^\]]*\]',body)
assert all(f'Page {n}' in texts[n-1] for n in range(1,len(pdf)+1))
for student in ['IT24103900','IT24103041','IT24100360','IT24102275']:assert student in body
with zipfile.ZipFile(F/'SE3110_Testing_Report_GR42_Final.docx') as z:assert z.testzip() is None
originalchecks=json.load(open(O/'audit/input-hashes.json'))
for path,h in originalchecks.items():assert hashlib.sha256(Path(path).read_bytes()).hexdigest()==h
bad=[]
for i,p in enumerate(pdf):
 for b in p.get_text('dict')['blocks']:
  for line in b.get('lines',[]):
   for span in line['spans']:
    a,b,c,e=span['bbox']
    if a<0 or b<0 or c>p.rect.width+1 or e>p.rect.height+1:bad.append(i+1)
assert not bad
r={'checked_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'docx_zip_opens':True,'docx_python_docx_opens':True,'pdf_pymupdf_opens':True,'pages':len(pdf),'original_tables':len(original.tables),'preserved_tables':len(d.tables),'toc_entries_verified':len(toc),'case_evidence_paths_verified':77,'page_footers_consistent':True,'student_ids_preserved':True,'original_input_hashes_unchanged':True,'text_outside_page_bounds':bad,'visual_review':'Every rendered page inspected; changed table/case/TOC pages rechecked after revisions. Landscape case tables readable, repeated headers, no visibly truncated rows. Sparse pages from original explicit member/reference breaks retained. Report contains no inline screenshots; current actual UI images are supplied as separate evidence and explicitly scoped to mocked API tests.','limitations':'LibreOffice rendering verifies opening/layout; no interactive Microsoft Word session was available. Historical external citations/deployed cover URLs are retained references, not verified service availability.'}
(O/'audit/final-quality-check.json').write_text(json.dumps(r,indent=2));print(json.dumps(r,indent=2))
