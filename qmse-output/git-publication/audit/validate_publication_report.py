from pathlib import Path
from datetime import datetime,timezone
import json,hashlib,zipfile,re,collections,shutil,subprocess,xml.etree.ElementTree as ET
import pymupdf
from docx import Document
R=Path.cwd();O=R/'qmse-output';G=O/'git-publication';F=O/'final-report';name='SE3110_Testing_Report_GR42_Final';render=F/'render-publication';shutil.copy2(render/(name+'.pdf'),F/(name+'.pdf'));old=Document(G/'backup-60-page'/(name+'.docx'));D=Document(F/(name+'.docx'));pdf=pymupdf.open(F/(name+'.pdf'));checks={}
def check(k,b):
 checks[k]=bool(b)
 if not b:raise AssertionError(k)
check('docx_opens',zipfile.ZipFile(F/(name+'.docx')).testzip() is None);check('tables_sections_preserved',len(D.tables)==len(old.tables)==71 and len(D.sections)==len(old.sections)==3);check('section_layout_unchanged',[s._sectPr.xml for s in D.sections]==[s._sectPr.xml for s in old.sections]);check('cover_students_unchanged',D.tables[1]._tbl.xml==old.tables[1]._tbl.xml)
for i in range(10,26):check('test_table_unchanged_'+str(i),D.tables[i]._tbl.xml==old.tables[i]._tbl.xml)
check('component_summary_unchanged',D.tables[45]._tbl.xml==old.tables[45]._tbl.xml)
def media(p):
 with zipfile.ZipFile(p) as z:return {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
check('all_20_previous_images_preserved',media(G/'backup-60-page'/(name+'.docx'))<=media(F/(name+'.docx')));check('30_genuine_figures',len(D.inline_shapes)==30)
figs=json.load(open(O/'evidence/visual/figure-index.json'))+json.load(open(O/'completion-20261008/evidence/figure-index.json'))+json.load(open(G/'evidence/figure-index.json'));check('unique_figure_numbers',len(figs)==len({f['figure'] for f in figs})==30)
norm=lambda x:re.sub(r'\s+',' ',x).strip();pages=[norm(p.get_text()) for p in pdf];placements=[]
for f in figs:
 title=norm('Figure '+f['figure']+'. '+f['title'])
 if f['figure']=='6.13':title='Figure 6.13. Deployment repository: historical Mobile CI success'
 found=[i+1 for i,t in enumerate(pages) if title in t];check('unique_caption_'+f['figure'],len(found)==1);p=pdf[found[0]-1];check('image_on_caption_page_'+f['figure'],bool(p.get_image_info()));check('figure_image_exists_'+f['figure'],(O/f['image']).exists())
 for src in f['source'].split('; '):
  if not src.startswith('https://'):check('figure_source_exists_'+src,(O/src).exists())
 placements.append({'figure':f['figure'],'page':found[0]})
check('30_pdf_image_placements',sum(len(p.get_image_info()) for p in pdf)==30);check('images_inside_pages',all(p.rect.contains(pymupdf.Rect(im['bbox'])) for p in pdf for im in p.get_image_info()));check('footer_page_numbers',all('Page '+str(i+1) in p.get_text() for i,p in enumerate(pdf)))
check('no_blank_pages',all(len(re.sub(r'SE3110 Software Testing Report.*|Page \d+','',p.get_text()).strip())>40 for p in pdf))
full='\n'.join(pages);check('no_outdated_unpushed_claim',not re.search(r'unpushed|no push|push approval|clone https://github.com/Jeyamdev',full,re.I));check('all_four_member_sections',all(s in full for s in ['9.1 Jeyashankar','9.2 Buddhika','9.3 Vithusan','9.4 Wijayananda']))
for p,h in json.load(open(O/'audit/input-hashes.json')).items():check('original_unchanged_'+Path(p).name,hashlib.sha256(Path(p).read_bytes()).hexdigest()==h)
cs=json.load(open(O/'completion-20261008/test-results/case-verification.json'));check('77_case_totals',dict(collections.Counter(c['status'] for c in cs))=={'VERIFIED':23,'FAILED':3,'BLOCKED':2,'HISTORICAL/UNVERIFIED':49})
for c in cs:
 for e in c['evidence']:check('case_evidence_'+e,(R/e).exists())
ns={'t':'http://microsoft.com/schemas/VisualStudio/TeamTest/2010'};tree=ET.parse(O/'completion-20261008/evidence/backend-final-verified.trx');count=tree.find('.//t:Counters',ns).attrib;check('xunit_328_326_0_2',count['total']=='328' and count['passed']=='326' and count['failed']=='0' and sum(x.attrib['outcome']=='NotExecuted' for x in tree.findall('.//t:UnitTestResult',ns))==2);ps=json.load(open(O/'completion-20261008/evidence/playwright-final.json'))['stats'];check('playwright_111_0_3',ps['expected']==111 and ps['unexpected']==0 and ps['skipped']==3);ai=ET.parse(O/'completion-20261008/evidence/ai-completion.xml').find('testsuite');check('pytest_102',ai.attrib['tests']=='102' and ai.attrib['failures']=='0' and ai.attrib['errors']=='0')
v=json.load(open(G/'evidence/publication-verification.json'));check('published_sha',v['approved_sha_matches'] and v['remote_sha']=='05359214df52fa743b234e1d77e6671f57c14dea');check('83_commits_existing_tips_unchanged',v['commits_beyond_main']==83 and v['all_existing_branch_tips_unchanged']);clone=json.load(open(G/'evidence/fresh-clone-verification.json'));check('reference_clone_authorship',clone['status']=='PASS' and clone['all_commit_author_and_committer_metadata_matches'] and clone['reachable_commits']==97)
ci=json.load(open(G/'evidence/ci-verification.json'));check('no_CI_pass_claim',ci['sha_runs']==0 and ci['branch_runs']==0 and all(not w['new_branch_matches'] and not w['workflow_dispatch_available'] for w in ci['source_workflows']))
toc=json.load(open(G/'audit/toc-pages.json'));check('54_toc_targets',len(toc)==54 and all(norm(x['title']) in pages[x['page']-1] for x in toc.values()))
for m in json.load(open(G/'evidence/published-capture-manifest.json')):
 check('genuine_GitHub_'+m['id'],m.get('genuine_page') and m['http_status']==200);check('GitHub_capture_hash_'+m['id'],hashlib.sha256((O/m['image']).read_bytes()).hexdigest()==m['sha256'])
# Check all portable text/docx/PDF outputs against configured private values and any historical database secret; never emit values.
secrets=[]
for p in Path('/private/tmp/qmse-completion-control').glob('*.json'):
 if p.name in ['settings.json','performance.json']:
  secrets += [v for k,v in json.load(open(p)).items() if k in ['password','key','token'] and isinstance(v,str) and len(v)>12]
for line in (R/'ai-service/.env').read_text().splitlines():
 if '=' in line:
  k,v=line.split('=',1);v=v.strip().strip('"\'')
  if re.search(r'KEY|SECRET|TOKEN',k,re.I) and len(v)>16:secrets.append(v)
# Pre-existing main connection blob: only private comparison values, no content copied.
blob=subprocess.check_output(['git','cat-file','blob','2cdc273dc09a42849255186adb46c10a53b66620']).decode(errors='ignore')
secrets += re.findall(r'npg_[A-Za-z0-9]{8,}',blob)
hits=[];scanned=0
for p in O.rglob('*'):
 if not p.is_file() or p.suffix.lower() in ['.zip','.png'] or any(x.startswith('render') or x in ['.dotnet','__pycache__'] for x in p.relative_to(O).parts):continue
 try:
  if p.suffix.lower()=='.docx':
   with zipfile.ZipFile(p) as z:txt='\n'.join(z.read(n).decode(errors='ignore') for n in z.namelist() if n.endswith('.xml'))
  elif p.suffix.lower()=='.pdf':txt='\n'.join(q.get_text() for q in pymupdf.open(p))
  else:txt=p.read_text()
  scanned+=1
 except (UnicodeError,OSError,zipfile.BadZipFile):continue
 if any(s in txt for s in secrets):hits.append(str(p.relative_to(R)))
check('private_values_excluded',not hits)
result={'checked_utc':datetime.now(timezone.utc).isoformat(),'all_checks_passed':all(checks.values()),'pdf_pages':len(pdf),'figures':30,'preserved_figures':20,'tables':71,'toc_targets':54,'test_case_totals':dict(collections.Counter(c['status'] for c in cs)),'figure_pages':placements,'secret_files_scanned':scanned,'checks':checks,'visual_review':'Pending separate page-by-page visual review record'};(G/'audit/final-quality-check.json').write_text(json.dumps(result,indent=2));print(json.dumps({k:v for k,v in result.items() if k not in ['checks','figure_pages']},indent=2))
