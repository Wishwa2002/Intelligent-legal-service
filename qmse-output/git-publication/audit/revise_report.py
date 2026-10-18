from pathlib import Path
from copy import deepcopy
from datetime import datetime,timezone
import json,re,subprocess
from docx import Document
from docx.shared import Pt,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
O=Path('qmse-output');G=O/'git-publication';F=O/'final-report';D=Document(G/'backup-60-page/SE3110_Testing_Report_GR42_Final.docx');P=list(D.paragraphs);T=D.tables;edits=[]
sha='05359214df52fa743b234e1d77e6671f57c14dea';repo='https://github.com/Wishwa2002/Intelligent-legal-service';branch=repo+'/tree/se3110-qmse-testing';deployment='https://github.com/Jeyamdev/legalease-ai'
def put(p,text,loc='paragraph'):
 old=p.text
 if old==text:return
 for child in list(p._p):
  if child.tag!=qn('w:pPr'):p._p.remove(child)
 p.add_run(text);edits.append({'location':loc,'before':old,'after':text})
def cell(c,text,loc='cell',size=None):
 p=c.paragraphs[0];put(p,text,loc)
 for e in list(c._tc):
  if e.tag==qn('w:p') and e is not p._p:c._tc.remove(e)
 if size:
  for r in p.runs:r.font.size=Pt(size)
def row(t,i,values):
 for j,v in enumerate(values):cell(t.cell(i,j),str(v),f'table {T.index(t)} row {i} cell {j}')
def addrow(t,values,after=None):
 e=deepcopy(t.rows[-1]._tr)
 if after is None:t._tbl.append(e)
 else:t.rows[after]._tr.addnext(e)
 from docx.table import _Row
 r=_Row(e,t)
 for c,v in zip(r.cells,values):cell(c,str(v))
 return r
def before(anchor,text,size=10,style=None):
 p=anchor.insert_paragraph_before(text,style);p.paragraph_format.space_after=Pt(6)
 for r in p.runs:r.font.size=Pt(size)
 return p
# Cover: retain title and student identities, distinguish repository roles explicitly.
row(T[0],3,['Original group development repository',repo])
r=addrow(T[0],['Final QMSE testing branch',branch],after=3)
r=addrow(T[0],['Deployment repository',deployment],after=4)
for ri in [3,4,5]:
 for ci,c in enumerate(T[0].rows[ri].cells):
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
   for run in p.runs:run.font.size=Pt(9);run.bold=ci==0
put(P[12],'Publication revision: 8 October 2026 (Asia/Colombo). Approved tested source 05359214df52, including fixes f360e4ac21f1, is published on Wishwa2002/Intelligent-legal-service, branch se3110-qmse-testing. Original development history is preserved; no deployment or submission occurred. GitHub CI did not trigger for this branch. Available SE3090 Version 3.0 reference reviewed; formal approval unconfirmed. Original, 51-page and pre-publication 60-page reports are preserved.')
before(P[13],'Repository provenance: Wishwa2002/Intelligent-legal-service is the original group development repository. Its main/develop/feature histories retain recorded authors and dates; the new testing branch preserves that ancestry. Jeyamdev/legalease-ai is the deployment repository used for hosting and recent local QMSE preparation. It is not the original group-history repository. Eleven newer deployment-only commits were excluded because they were outside the verified baseline.')
put(P[24],'Keep reusable automated test scripts in the published testing branch and retain native reports, screenshots and final documents in the separately supplied evidence package so results can be reproduced and explained.')
row(T[6],6,['Version control','Published Wishwa2002/Intelligent-legal-service / se3110-qmse-testing at '+sha+'. Local source branch: docs/se3110-qmse-testing. Identical tested code; CI not triggered.'])
cell(T[2].cell(6,1),'.github/workflows/backend-ci.yml, frontend-ci.yml and mobile-ci.yml have main/develop and component-path filters. Published testing SHA: zero observed runs; no manual dispatch. GitHub-managed CodeQL/Dependabot metadata is separately retained.')
before(P[38],'Source baseline for all case references: application/test fixes f360e4ac21f149ae5b90a125c57bf928982a9c53 and replay-script commit '+sha+'. The exact objects are now published on the original group repository testing branch. Publication does not change prior local execution outcomes or convert historical cases into passes.')
# Preserve historical CI screenshot and qualify its provenance.
put(P[101],'Figures 6.1–6.13 remain genuine dated earlier-phase evidence. Figures 6.14–6.20 retain the verified fixes and real local workflow; new Figures 6.21–6.30 establish original group history, member commit records, publication and CI non-trigger. Historical deployment-repository images are preserved with provenance. Full native sources are included; GitHub screenshots prove repository observations, not successful test execution. Paths are relative to qmse-output.')
put(P[162],'Figure 6.13. Deployment repository: historical Mobile CI success')
put(P[166],'This preserved Jeyamdev/legalease-ai page shows historical Mobile CI Success for base commit 8a9b8a2, duration 2m18s; other historical runs failed. It does not cover later QMSE commits f360e4a/0535921 or prove group-branch CI success. Figure 6.30 records the published group-branch non-trigger.')
cell(T[42].cell(9,3),'evidence/ci/github-actions-runs.json; historical deployment/base revision only. Group publication CI: git-publication/evidence/ci-verification.json')
put(P[219],'The original component assignments and member disclosures are retained. Source/test files and 26 original member-associated commit records were checked in the original group repository, with author/date/path evidence and preserved ancestry. Personal ownership, effort and responsibilities still require member confirmation. The two new Codex-assisted QMSE commits are recorded separately and are not individual contribution proof.')
# Exact attribution: recorded aliases and supporting paths, member effort still requires confirmation.
members=json.load(open(G/'evidence/member-commit-evidence.json'))
for mi,(table,heading) in enumerate([(48,232),(53,250),(58,268),(63,286)]):
 m=members[mi];t=T[table]
 while len(t.rows)-1<len(m['commits']):addrow(t,['','',''])
 for ri,c in enumerate(m['commits'],1):
  row(t,ri,[c['sha'][:7],c['author_date'][:10],c['subject']+'\nRecorded author: '+c['author']+'; '+c['author_date'][11:]+'.'])
  p=t.cell(ri,0).paragraphs[0];put(p,c['sha'][:7]);h=OxmlElement('w:hyperlink');rid=D.part.relate_to(c['url'],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);h.set(qn('r:id'),rid);run=OxmlElement('w:r');prop=OxmlElement('w:rPr');u=OxmlElement('w:u');u.set(qn('w:val'),'single');prop.append(u);run.append(prop);tx=OxmlElement('w:t');tx.text=c['sha'][:7];run.append(tx);h.append(run)
  for child in list(p._p):
   if child.tag!=qn('w:pPr'):p._p.remove(child)
  p._p.append(h)
 note=f"Original group commit evidence ({len(m['commits'])} records): "+', '.join(m['recorded_aliases'])+'. The linked commits exist in Wishwa2002/Intelligent-legal-service and are ancestors of the published tested SHA. Author dates and changed paths are recorded in git-publication/evidence/member-commit-evidence.json. Identity association uses the supplied group context; personal effort, responsibility and AI use require member confirmation. No new QMSE work is attributed personally.'
 before(P[heading+1],note)
cell(T[61].cell(5,1),'Historical responsibility claim; member confirmation pending. Actual source files are backend-ci.yml, frontend-ci.yml and mobile-ci.yml. Pushes to se3110-qmse-testing do not match main/develop filters; zero new-SHA runs observed. No every-push/pass claim is supported.')
put(P[283],'Published testing branch Actions check: zero runs for SHA '+sha+' and branch se3110-qmse-testing. The three committed source workflows allow push only to main/develop and none declares workflow_dispatch. Managed CodeQL/Dependabot have no observed run for this revision; their trigger settings are not established by source YAML. Separate unapplied CI proposals are retained for review.')
before(P[291],'Consultation ownership limitation: history for app/agents/supervisor.py and its prompt records a566efa/6b59d19 under VTN02. Do not assign this implementation or RAG scoring to Wijayananda solely because it belongs to C4. Shared work and identity mapping require group confirmation. Authentication and service-request evidence under Wishwa is independently recorded (43f6345, 9309287, 9579cee).')
before(P[295],'Contribution matrix provenance: original member development records are linked in Section 9 and Figures 6.23–6.26. The matrix test counts are observed suite cases, not personal-effort percentages. New Codex-assisted fixes and replay scripts are recorded separately under the actual Git author Jeyamdev (f360e4a, 0535921); no reassignment to students is made. Original main has 14 reachable commits; the testing branch has 97, including 83 beyond main. All original object identities and authors are intact.')
for i,m in enumerate(members,1):cell(T[66].cell(i,2),'Group-repository development: '+', '.join(c['sha'][:7] for c in m['commits'])+'. Personal testing effort/ownership requires confirmation; published QMSE additions are Codex-assisted.')
put(P[296],P[296].text.replace('CI for the unpushed tested commit remain outstanding','CI execution for the published tested commit remain outstanding').replace('assessment completeness is not claimed','the published branch is verified, but assessment completeness is not claimed'))
put(P[303],'Resolve provider availability/model/budget before further live Golden Cases. Obtain member contribution confirmations and complete native/accessibility/manual demonstrations. CI needs a separately approved configuration revision: current branch filters prevent execution, and no manual dispatch is declared. Publication of the verified SHA is complete; no approval for an additional SHA, deployment or submission is inferred.')
cell(T[67].cell(2,1),'Codex assisted tests, four fixes, isolated local execution, repository audit, approved exact-SHA publication and this report revision. No production data was modified and no deployment or submission occurred. New work is not attributed to students without confirmation. Git author metadata is preserved; it is not proof of personal effort.')
cell(T[29].cell(2,1),'.github/workflows/mobile-ci.yml (current path; historical defect reference)')
row(T[50],4,['Recommendation AI','From repository root: cd ai-service\n.venv/bin/python -m uvicorn lawyer_recommendation.app:app --host 127.0.0.1 --port 8002'])
row(T[55],4,['Scheduling AI service','From repository root: cd ai-service\n.venv/bin/python -m uvicorn app.main:app --host 127.0.0.1 --port 8001'])
row(T[60],2,['Document AI service','From repository root: cd ai-service\n.venv/bin/python -m uvicorn app.main:app --host 127.0.0.1 --port 8001'])
# Reference and compliance corrections.
put(P[308],'SE3090 LegalEase AI Full Report, Version 3.0 — Consolidated Final Verified Specification, Group 2026-SE-42. Available requirements reference; formal approval not independently confirmed.')
for text in ['Original group development repository: '+repo,'Final published QMSE testing branch: '+branch+'; expected SHA '+sha+'.','Deployment repository: '+deployment+' (hosting/deployment history; historical CI Figure 6.13).','Publication and commit evidence: git-publication/evidence/publication-verification.json, fresh-clone-verification.json, member-commit-evidence.json and ci-verification.json; captured 8 October 2026.']:
 before(P[314],text,9)
cell(T[68].cell(14,2),'VERIFIED: reusable test source and scripts published at '+sha+' on group se3110-qmse-testing. Native logs/final report supplied separately; no runtime caches published.')
cell(T[68].cell(15,2),'VERIFIED publication/ancestry; 26 original member-associated commits with authors/dates/paths. Member effort confirmation remains pending. CI NOT TRIGGERED: main/develop branch filters, no manual dispatch.')
# Reproduction commands from published layout; never imply a new test execution.
row(T[69],1,['Clone and verify','git clone https://github.com/Wishwa2002/Intelligent-legal-service.git\ncd Intelligent-legal-service\ngit switch se3110-qmse-testing\ngit rev-parse HEAD\nExpected: '+sha])
row(T[69],2,['Scheduling / document AI (:8001)','From repository root in a dedicated terminal:\ncd ai-service\npython -m venv .venv\n.venv/bin/python -m pip install -r requirements.txt\n.venv/bin/python -m uvicorn app.main:app --host 127.0.0.1 --port 8001'])
row(T[69],3,['Recommendation AI (:8002)','From repository root in a separate terminal:\ncd ai-service\n.venv/bin/python -m uvicorn lawyer_recommendation.app:app --host 127.0.0.1 --port 8002\nLive provider credentials/usage approval required; offline tests use mocks.'])
row(T[69],4,['Backend API','Configure only a fresh isolated PostgreSQL connection, generated local JWT/bridge keys and matching AI URLs. From root:\ndotnet restore LegalService.sln\ndotnet build LegalService.sln\ndotnet run --project backend/LegalService.API/LegalService.API.csproj --launch-profile http\nProfile binds 5000 and 5295; align clients explicitly. Never use a production connection.'])
row(T[69],5,['React web','From repository root:\ncd frontend\nnpm ci\nnpm run dev -- --host 127.0.0.1\nConfigure local VITE_API_URL outside version control.'])
addrow(T[69],['Flutter application','From repository root:\ncd mobile\nflutter pub get\nConfigure ignored .env for the isolated BACKEND_URL.\nflutter run -d chrome\nFor recorded fixed-port web workflow: flutter build web --debug, then serve mobile/build/web at 55448; generated asset must target local API 55443.'])
row(T[70],1,['Backend (xUnit)','From repository root:\ndotnet test backend/LegalService.Tests/LegalService.Tests.csproj --logger "trx;LogFileName=backend.trx" --results-directory qmse-output/reproduction/backend\nFor relational tests set MEMBER1_TEST_POSTGRES and LAWYER_MOBILE_TEST_POSTGRES to a fresh loopback test cluster; fixture databases are created/dropped. Without them PostgreSQL tests can skip.'])
row(T[70],2,['Backend coverage','dotnet test backend/LegalService.Tests/LegalService.Tests.csproj --collect:"XPlat Code Coverage" --results-directory qmse-output/reproduction/coverage'])
row(T[70],3,['Node helper tests','cd frontend\nnpm ci\nnpm run test:member1'])
row(T[70],4,['Playwright component UI','cd frontend\nnpx playwright install chromium\nnpm test -- --reporter=list,html\nnpx playwright show-report\nMost spec routes are mocked; this is separate from the real integrated workflow.'])
row(T[70],5,['Flutter unit/widget','cd mobile\nflutter pub get\nflutter test --machine\nRecorded final 49-pass output: test-results/completion-20261008-mobile.log; prior attempts preserved.'])
row(T[70],6,['Security scan','Prerequisite: ZAP and actual isolated API/PostgreSQL. Retained script: qmse-output/completion-20261008/audit/expanded_zap.py. Confirm its loopback ports and private control files; default recorded mode is passive plus finite manual probes, not automated active scanning.'])
row(T[70],7,['Controlled k6','k6 run --summary-export=qmse-output/reproduction/k6-summary.json qmse-output/completion-20261008/audit/postgres-smoke.k6.js\nRequires private /private/tmp/qmse-completion-control/performance.json (local API URL, locally generated token and synthetic lawyerId); 2 VUs/15 s. Run in a fresh clone or adapt the hard-coded summary output path to preserve previous logs.'])
addrow(T[70],['AI deterministic/mocked','cd ai-service\nGEMINI_API_KEY=qmse-synthetic-not-a-real-provider-key GOOGLE_API_KEY=qmse-synthetic-not-a-real-provider-key PYTHON_DOTENV_DISABLED=1 HF_HUB_OFFLINE=1 TRANSFORMERS_OFFLINE=1 .venv/bin/python -m pytest tests -s --junitxml=../qmse-output/reproduction/ai.xml\nProvider calls remain disabled; do not claim live Gemini results.'])
addrow(T[70],['Actual integrated workflow','With fresh fixed-port services and private synthetic settings established, from root:\nnode frontend/tests/qmse-integration/integrated_workflow.cjs\nnode frontend/tests/qmse-integration/resume_workflow.cjs\nnode frontend/tests/qmse-integration/complete_review.cjs\nVerify integrated-result.json AND HTTP/SQL outputs, not intermediate exit codes.'])
put(P[321],'Recorded real-workflow setup and replay details are in frontend/tests/qmse-integration/README.md and git-publication/audit/reproduction-published.md. The committed supervisor uses macOS-specific runtime paths; adapt them to the assessor host. A clone contains reusable scripts, not local caches, private settings or the separately delivered native evidence. Build each service and use fresh isolated PostgreSQL fixtures before replay. The reference-assisted publication clone verified remote head and original commit metadata using audited local objects; no application rerun or independent full-object download is claimed. A first unassisted clone timed out and is retained.')
# Add genuine post-push GitHub figures, after original twenty, before Section 7.
manifest={m['id']:m for m in json.load(open(G/'evidence/published-capture-manifest.json'))};main='59337bba0227aa07e18e60f697bd758db327f0bd'
items=[('official-overview','Official group repository overview','main',main,'The original group repository remains on its unchanged main revision. This overview establishes repository identity; the final tested implementation is on a separate branch.'),('original-main-history','Original main-branch development history','main',main,'Original dated main commits and author records are preserved. Main has 14 reachable commits under Wishwa and yhprasad-related author strings; other member records occur in incorporated develop/feature history.'),('member1-lawyer-development','Original lawyer-management development record','develop / tested-branch ancestry',members[0]['commits'][1]['sha'],'Recorded author Jeyamdev; lawyer/mobile/matching/workforce changes. Related specialization and recommendation commits are linked in Section 9.1. This is historical source evidence; personal effort requires confirmation.'),('member2-appointments','Original appointment-management development record','develop / tested-branch ancestry',members[1]['commits'][-1]['sha'],'Recorded author yhprasad (GitHub account yhprasadhew); appointment APIs, conflict detection and status transitions. Scheduling/careers records are in Section 9.2. This does not prove historical test execution.'),('member3-document-agent','Original document/clerk/AI development record','Document-&-Cleark / tested-branch ancestry',members[2]['commits'][-1]['sha'],'Recorded author VTN02; document validation, mobile client and clerk management. Section 9.3 records related commits. Shared supervisor code appears in this history, so component assignment alone cannot establish ownership.'),('member4-authentication','Original authentication and authorization record','feature/admin-auth / tested-branch ancestry',members[3]['commits'][-2]['sha'],'Recorded author Wishwa (GitHub account Wishwa2002); JWT/admin authorization and service-request changes. Section 9.4 records related security/service commits and qualifies consultation ownership.'),('published-branch','Published final QMSE testing branch','se3110-qmse-testing',sha,'The branch selector identifies the published testing branch. This narrow capture retains asynchronous counter placeholders; the earlier genuine wide capture and native Git verification establish 83 commits beyond main and 97 reachable commits. The remote head matches the approved SHA. All 15 pre-existing branch tips were unchanged immediately after publication; eleven newer deployment-only commits were excluded.'),('published-tested-commit','Published tested source and replay-script commit','se3110-qmse-testing',sha,'Actual published commit 0535921 has parent f360e4a. The nine-file replay-script addition preserves executed tooling. Recorded author is Jeyamdev; these Codex-assisted QMSE additions are not original student-development attribution.'),('testing-history','Published testing-related commit history','se3110-qmse-testing',sha,'The two new QMSE commits f360e4a and 0535921 sit above preserved development/deployment ancestry. Git authors and dates remain unchanged. Historical check icons belong to their own revisions, not the latest tested head.'),('actions-for-tested-revision','GitHub Actions: no run for the published testing branch','se3110-qmse-testing',sha,'Actual GitHub branch filter shows zero workflow results; read-only API queries also show zero for the exact SHA. Backend, frontend and mobile push filters permit only main/develop; no workflow_dispatch is declared. No CI pass is claimed. Managed workflow trigger settings are not established by source YAML.')]
figs=[]
for number,(id,title,b,s,explanation) in enumerate(items,21):
 m=manifest[id];assert m.get('genuine_page') and (O/m['image']).exists();p=before(P[203],f'Figure 6.{number}. {title}');p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True;p.runs[0].bold=True
 p=before(P[203],'');p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(O/m['image']),width=Inches(6.1))
 author_meta=subprocess.check_output(['git','show','-s','--format=%an; author date %aI; committer date %cI',s],text=True).strip()
 for txt,size in [('Repository: Wishwa2002/Intelligent-legal-service; branch: '+b+'.',9),('Relevant SHA: '+s,9),('Git record: '+author_meta,8.5),('Capture: '+m['captured_utc']+' (UTC); source: '+m['url'],8.5),(explanation,10)]:
  p=before(P[203],txt,size);p.paragraph_format.keep_with_next=txt!=explanation
 figs.append({'figure':f'6.{number}','title':title,'execution_date':m['captured_utc'],'source':'git-publication/evidence/'+id+'.png; '+m['url'],'image':m['image'],'repository':'Wishwa2002/Intelligent-legal-service','branch':b,'sha':s,'interpretation':explanation})
addrow(T[42],['Figs 6.21–6.22','Original group repository and main history','Genuine GitHub / unchanged original commits','git-publication/evidence/original-main-history.png; group-main-commits.json'])
addrow(T[42],['Figs 6.23–6.26','Member-associated original development records','Recorded authors + paths; personal effort pending','git-publication/evidence/member-commit-evidence.json'])
addrow(T[42],['Figs 6.27–6.30','Published branch, source SHA, ancestry and Actions non-trigger','Git remote/API/browser verification','git-publication/evidence/publication-verification.json; ci-verification.json; fresh-clone-verification.json'])
# Real historical risk, recorded without copying any connection data.
before(P[211],'Git publication security note: GitHub push output reported nine default-branch Dependabot alerts (seven high, two moderate); no authenticated advisory review or remediation was performed. These are separate from local ZAP findings. A database-connection artifact was also identified in pre-existing public main history and is absent from the tested tree. Its contents/credential-like filename are excluded from deliverables; the group must review rotation/removal with repository owners. No history rewrite was authorised.')
# Make genuine repository URLs clickable while preserving existing run formatting.
for para in D.element.findall('.//'+qn('w:p')):
 for oldrun in list(para):
  if oldrun.tag!=qn('w:r'):continue
  tt=oldrun.findall(qn('w:t'))
  if len(tt)!=1:continue
  text=tt[0].text or '';matches=list(re.finditer(r'https://github\.com/[A-Za-z0-9_./?=%:+-]+',text))
  if not matches:continue
  at=0;nodes=[]
  def fragment(value):
   run=OxmlElement('w:r');pr=oldrun.find(qn('w:rPr'))
   if pr is not None:run.append(deepcopy(pr))
   tx=OxmlElement('w:t');tx.set('{http://www.w3.org/XML/1998/namespace}space','preserve');tx.text=value;run.append(tx);return run
  for match in matches:
   url=match.group().rstrip('.;,')
   if match.start()>at:nodes.append(fragment(text[at:match.start()]))
   link=OxmlElement('w:hyperlink');link.set(qn('r:id'),D.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True));run=fragment(url);pr=run.find(qn('w:rPr'))
   if pr is None:pr=OxmlElement('w:rPr');run.insert(0,pr)
   color=OxmlElement('w:color');color.set(qn('w:val'),'005EA5');pr.append(color);under=OxmlElement('w:u');under.set(qn('w:val'),'single');pr.append(under);link.append(run);nodes.append(link);at=match.start()+len(url)
  if at<len(text):nodes.append(fragment(text[at:]))
  idx=para.index(oldrun);para.remove(oldrun)
  for node in nodes:para.insert(idx,node);idx+=1

# Record all body/cell text for whole-document review; preserve headings/sections/tables and test case results.
D.save(F/'SE3110_Testing_Report_GR42_Final.docx');(G/'evidence/figure-index.json').write_text(json.dumps(figs,indent=2));(G/'audit/document-change-log.json').write_text(json.dumps(edits,indent=2));after=[]
for i,p in enumerate(D.paragraphs):after.append({'type':'paragraph','index':i,'style':p.style.name,'text':p.text})
for ti,t in enumerate(D.tables):
 for ri,r in enumerate(t.rows):
  for ci,c in enumerate(r.cells):after.append({'type':'cell','table':ti,'row':ri,'cell':ci,'text':c.text})
(G/'audit/document-after-blocks.json').write_text(json.dumps(after,indent=2));print(json.dumps({'paragraphs':len(D.paragraphs),'tables':len(D.tables),'sections':len(D.sections),'genuine_figures':len(D.inline_shapes),'edits':len(edits)}))
